# Create a polished Word document for the ALM Business Glossary with executive formatting

from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()

# Set default style
style = doc.styles['Normal']
font = style.font
font.name = 'Calibri'
font.size = Pt(11)

# Title Page
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run("ALM Business Glossary\n& Inventory Classification Framework")
run.bold = True
run.font.size = Pt(20)

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = subtitle.add_run("\nExecutive & Analytical Reference Document")
run.font.size = Pt(12)

meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = meta.add_run("\nVersion 1.1\nJanuary 2026")
run.italic = True

doc.add_page_break()

# Executive Summary
h = doc.add_heading("Executive Summary", level=1)

doc.add_paragraph(
    "The ALM Business Glossary establishes a shared, authoritative understanding of core financial and inventory concepts "
    "used in executive reporting, analytics, and strategic decision-making. This document addresses long-standing challenges "
    "caused by inconsistent terminology, subjective inventory classifications, and misaligned interpretations of key metrics.\n\n"
    "A central focus of this framework is inventory movement classification based on estimated time required to sell through "
    "current on-hand inventory. This forward-looking approach aligns inventory analysis with cash recovery expectations and "
    "supports clearer, more defensible decision-making at the executive level."
)

doc.add_paragraph(
    "Inventory movement status is determined using a sell-through time horizon derived from recent sales velocity. "
    "This replaces purely backward-looking aging metrics with a practical assessment of inventory risk."
)

table = doc.add_table(rows=1, cols=3)
hdrs = table.rows[0].cells
hdrs[0].text = "Movement Status"
hdrs[1].text = "Estimated Time to Sell Through"
hdrs[2].text = "Executive Interpretation"

rows = [
    ("Active – Moving", "< 18 months", "Healthy inventory"),
    ("Active – Slow Moving", "18–36 months", "Monitor and manage"),
    ("Active – No Movement", "> 36 months", "Elevated risk")
]

for r in rows:
    row = table.add_row().cells
    for i, val in enumerate(r):
        row[i].text = val

doc.add_page_break()

# Business Glossary Section
doc.add_heading("ALM Business Glossary", level=1)

doc.add_paragraph(
    "This section defines enterprise-level financial and inventory terms used across ALM leadership reporting, analytics, "
    "and decision-making. Definitions prioritize business clarity and consistency and are intended to serve as the single "
    "authoritative reference for interpretation."
)

terms = [
    ("Revenue", "Total value of goods sold before deductions."),
    ("Net Sales", "Revenue net of discounts, returns, and allowances."),
    ("COGS", "Direct costs attributable to products sold."),
    ("Gross Margin", "Net Sales minus COGS."),
    ("Inventory Value", "On-hand inventory value based on standard costing."),
    ("Landed Cost", "Total cost to bring product to warehouse ready for sale."),
    ("Tariff Cost", "Government-imposed duties on imported goods.")
]

for term, definition in terms:
    p = doc.add_paragraph()
    r = p.add_run(term)
    r.bold = True
    doc.add_paragraph(definition)

doc.add_page_break()

# Inventory Classification
doc.add_heading("Inventory Lifecycle & Movement Classifications", level=1)

doc.add_paragraph(
    "ALM classifies inventory using two complementary dimensions: lifecycle status and movement status. "
    "Lifecycle reflects strategic intent, while movement reflects estimated sell-through timing."
)

doc.add_paragraph(
    "Movement status is derived from estimated months required to sell through current on-hand inventory, "
    "based on recent sales velocity."
)

doc.add_page_break()

# Appendix
doc.add_heading("Appendix A: Semantic Model Alignment (Internal Reference)", level=1)

doc.add_paragraph(
    "This appendix documents how glossary definitions and inventory classifications are operationalized within "
    "analytics and reporting models. It is intended for internal analytical reference only and not for executive distribution."
)

doc.add_paragraph(
    "Estimated Months to Sell Through is calculated as current on-hand quantity divided by average monthly units sold. "
    "Movement status is assigned based on this estimate and refreshed on a rolling basis."
)

# Save document
path = "/mnt/data/ALM_Business_Glossary_v1_1.docx"
doc.save(path)

path
