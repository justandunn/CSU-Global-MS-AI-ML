"""Build APA 7 analysis reports for CSC580 Modules 5 and 6."""
from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parent
M5 = ROOT / "MOD5" / "critical_thinking_option_2"
M6 = ROOT / "MOD6" / "critical_thinking_option_1"


def set_font(run, size: int = 12, bold: bool | None = None, italic: bool | None = None) -> None:
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Times New Roman")
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    set_font(run)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, end])


def setup_document(title: str, subject: str) -> Document:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.5)
    section.footer_distance = Inches(0.5)
    add_page_number(section.header.paragraphs[0])

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal.paragraph_format.line_spacing = 2
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)

    for style_name in ("Heading 1", "Heading 2"):
        style = doc.styles[style_name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(12)
        style.font.bold = True
        style.font.color.rgb = None
        style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(0)
        style.paragraph_format.keep_with_next = True
    doc.styles["Heading 1"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for _ in range(4):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run(title), bold=True)
    for line in (
        "Justan Dunn",
        "Colorado State University Global",
        "CSC580-1: Applying Machine Learning and Neural Networks",
        "Dong Nguyen",
        "August 30, 2026",
    ):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_font(p.add_run(line))

    doc.add_page_break()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    set_font(p.add_run(title), bold=True)
    doc.core_properties.title = title
    doc.core_properties.author = "Justan Dunn"
    doc.core_properties.subject = subject
    return doc


def add_body(doc: Document, text: str) -> None:
    p = doc.add_paragraph(text)
    p.paragraph_format.first_line_indent = Inches(0.5)
    p.paragraph_format.widow_control = True


def add_figure(doc: Document, image: Path, number: int, title: str, width: float = 6.1) -> None:
    number_line = doc.add_paragraph()
    number_line.paragraph_format.keep_with_next = True
    number_line.paragraph_format.line_spacing = 2
    set_font(number_line.add_run(f"Figure {number}"), bold=True)
    title_line = doc.add_paragraph()
    title_line.paragraph_format.keep_with_next = True
    title_line.paragraph_format.line_spacing = 2
    set_font(title_line.add_run(title), italic=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 1
    p.paragraph_format.space_after = Pt(6)
    p.add_run().add_picture(str(image), width=Inches(width))


def add_references(doc: Document, entries: list[list[tuple[str, bool]]]) -> None:
    doc.add_page_break()
    doc.add_heading("References", level=1)
    for entry in entries:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.5)
        p.paragraph_format.first_line_indent = Inches(-0.5)
        for text, italic in entry:
            set_font(p.add_run(text), italic=italic)


def build_module5() -> Path:
    results = json.loads((M5 / "evidence" / "results.json").read_text(encoding="utf-8"))
    title = "Option 2: Building a Random Forest Classifier"
    doc = setup_document(title, "CSC580 Module 5 Critical Thinking Assignment, Option 2")
    doc.add_heading("Purpose and Dataset", level=1)
    add_body(doc, "This assignment implemented a random forest classifier to predict Iris species from sepal length, sepal width, petal length, and petal width. The scikit-learn Iris dataset contains 150 observations, three equally represented species, and four continuous measurement variables (Scikit-learn Developers, 2026a). A fixed NumPy seed of zero made the assignment's random 75% training-mask procedure reproducible. The resulting partition contained 118 training observations and 32 held-out test observations.")
    add_figure(doc, M5 / "evidence" / "01_dataset_head.png", 1, "First Five Rows of the Iris Dataset", 6.2)
    add_figure(doc, M5 / "evidence" / "02_train_test_split.png", 2, "Reproducible Training and Test Assignment", 6.2)

    doc.add_heading("Preprocessing and Classifier Method", level=1)
    add_body(doc, "The four measurement columns were retained as predictors, and the species names in the training data were factorized as integer class labels. Random forests combine many decision trees trained with bootstrap samples and randomized feature selection. Aggregating decorrelated trees generally produces a more stable classifier than relying on one decision tree (Breiman, 2001). The submitted implementation used 100 trees, two parallel jobs, the Gini criterion, and a fixed model seed. No scaling was required because tree split decisions depend on ordered thresholds rather than feature magnitudes.")
    add_figure(doc, M5 / "evidence" / "03_preprocessing.png", 3, "Selected Features and Encoded Training Targets", 6.2)

    doc.add_heading("Predictions and Overall Performance", level=1)
    add_body(doc, f"The classifier produced a probability for each species and selected the class with the largest probability. On the {results['test_rows']}-observation test set, it correctly classified 30 observations and achieved {results['accuracy'] * 100:.2f}% accuracy. The first ten probability vectors demonstrate that the model often assigned high confidence to its selected class, while the actual-versus-predicted comparison supplies a direct record of the first five outcomes. Accuracy is appropriate for this balanced three-class problem, but the confusion matrix and per-class measures are needed to identify uneven performance.")
    add_figure(doc, M5 / "evidence" / "04_predicted_probabilities.png", 4, "Predicted Probabilities for the First Ten Test Observations", 6.2)
    add_figure(doc, M5 / "evidence" / "05_actual_vs_predicted.png", 5, "Actual and Predicted Species for the First Five Test Observations", 6.2)

    doc.add_heading("Confusion Matrix Interpretation", level=1)
    add_body(doc, "The confusion matrix shows perfect classification for all 13 setosa and all 12 virginica observations. Five of seven versicolor observations were correct, while two were classified as virginica. Versicolor recall was therefore 71.43%, compared with 100% recall for the other classes. This result is consistent with the feature space: setosa is well separated, whereas versicolor and virginica overlap in petal measurements. The errors are concentrated in one scientifically plausible boundary rather than distributed arbitrarily across all classes.")
    add_figure(doc, M5 / "evidence" / "06_confusion_matrix.png", 6, "Held-Out Test Confusion Matrix", 5.5)

    doc.add_heading("Feature Importance and Limitations", level=1)
    importances = results["feature_importance"]
    add_body(doc, f"Petal length ({importances['petal length (cm)']:.3f}) and petal width ({importances['petal width (cm)']:.3f}) jointly accounted for approximately {(importances['petal length (cm)'] + importances['petal width (cm)']) * 100:.1f}% of impurity-based importance. Sepal length contributed {importances['sepal length (cm)']:.3f}, and sepal width contributed {importances['sepal width (cm)']:.3f}. These values show that petal dimensions drove most decisions, but impurity importance should not be interpreted as causation. The evaluation also reflects only 32 test records from one fixed random split. Repeated stratified cross-validation and permutation importance would provide more stable estimates (Scikit-learn Developers, 2026b).")
    add_figure(doc, M5 / "evidence" / "07_feature_importance.png", 7, "Random Forest Feature Importance", 5.8)

    doc.add_heading("Conclusion", level=1)
    add_body(doc, "The random forest successfully classified the held-out Iris observations with 93.75% accuracy and made only two errors, both involving versicolor observations predicted as virginica. The model is reproducible, interpretable at the feature-importance level, and appropriate for demonstrating multiclass ensemble classification. Its primary limitations are the small test sample, reliance on one random partition, and the known overlap between two species. Cross-validation and uncertainty reporting would be the next steps before treating the result as a general performance estimate.")
    add_references(doc, [
        [("Breiman, L. (2001). Random forests. ", False), ("Machine Learning, 45", True), ("(1), 5-32. https://doi.org/10.1023/A:1010933404324", False)],
        [("Scikit-learn Developers. (2026a). ", False), ("load_iris", True), (". https://scikit-learn.org/stable/modules/generated/sklearn.datasets.load_iris.html", False)],
        [("Scikit-learn Developers. (2026b). ", False), ("RandomForestClassifier", True), (". https://scikit-learn.org/stable/modules/generated/sklearn.ensemble.RandomForestClassifier.html", False)],
    ])
    output = M5 / "CSC580_CTA5_Option_2_Dunn_Justan.docx"
    doc.save(output)
    return output


def build_module6() -> Path:
    results = json.loads((M6 / "evidence" / "results.json").read_text(encoding="utf-8"))
    title = "Option 1: CIFAR-10 Image Classification With a CNN"
    doc = setup_document(title, "CSC580 Module 6 Critical Thinking Assignment, Option 1")
    doc.add_heading("Purpose and Data", level=1)
    add_body(doc, "This assignment trained a convolutional neural network (CNN) to classify CIFAR-10 images. CIFAR-10 contains 60,000 color images measuring 32 by 32 pixels across 10 balanced classes: airplane, automobile, bird, cat, deer, dog, frog, horse, ship, and truck. The original partition provides 50,000 training images and 10,000 test images, with exactly 1,000 test images per class (Krizhevsky et al., 2009). Pixel values were converted to floating-point numbers and divided by 255 so every channel fell between zero and one.")
    add_figure(doc, M6 / "evidence" / "01_cifar10_samples.png", 1, "Representative CIFAR-10 Training Images", 6.1)

    doc.add_heading("Network Architecture and Training", level=1)
    add_body(doc, f"The TensorFlow {results['tensorflow_version']} model contained three convolutional stages with 32, 64, and 128 filters. Batch normalization stabilized intermediate activations, max pooling reduced spatial dimensions, and global average pooling limited the parameter count before the final dense classifier. Random horizontal flips and small translations augmented only the training stream. Dropout rates of 0.20, 0.30, and 0.40 provided additional regularization. The final model contained {results['parameters']:,} trainable and nontrainable parameters and used Adam, sparse categorical cross-entropy, a batch size of 128, and a maximum of 10 epochs. Ten percent of the training partition was reserved for validation. Early stopping restored the weights with the lowest validation loss, while learning-rate reduction responded to stalled validation improvement (TensorFlow, 2026).")

    doc.add_heading("Training Behavior", level=1)
    add_body(doc, f"Training completed {results['epochs_completed']} epochs and reached a best validation accuracy of {results['best_validation_accuracy'] * 100:.2f}%. The loss and accuracy curves in Figure 2 show whether improvements transferred from the training subset to the validation subset. A persistent training-validation gap indicates remaining overfitting, while continued validation improvement indicates that the model is learning reusable visual features. The augmentation, dropout, and stopping controls were included specifically to limit memorization of the 32-pixel training images.")
    add_figure(doc, M6 / "evidence" / "02_training_history.png", 2, "Training and Validation Performance by Epoch", 6.2)

    doc.add_heading("Held-Out Test Veracity", level=1)
    add_body(doc, f"The untouched 10,000-image test set produced {results['test_accuracy'] * 100:.2f}% accuracy and cross-entropy loss of {results['test_loss']:.3f}. This is a more credible estimate than training accuracy because test images did not influence gradient updates, learning-rate changes, early stopping, or model selection. However, a single aggregate accuracy value does not establish equal reliability across classes. The confusion matrix identifies which visual categories the model most often confuses and should be examined alongside class-level precision and recall.")
    add_figure(doc, M6 / "evidence" / "03_confusion_matrix.png", 3, "CIFAR-10 Held-Out Test Confusion Matrix", 6.0)

    doc.add_heading("Prediction Analysis", level=1)
    report = results["classification_report"]
    class_names = {"airplane", "automobile", "bird", "cat", "deer", "dog", "frog", "horse", "ship", "truck"}
    weakest = min((name for name in report if name in class_names), key=lambda name: report[name]["f1-score"])
    strongest = max((name for name in report if name in class_names), key=lambda name: report[name]["f1-score"])
    add_body(doc, f"The strongest class by F1 score was {strongest} ({report[strongest]['f1-score']:.3f}), while the weakest was {weakest} ({report[weakest]['f1-score']:.3f}). This variation is expected because some categories share textures and shapes at low resolution; for example, cats and dogs or deer and horses can be visually similar. Figure 4 displays 20 deterministic test examples with actual labels, predicted labels, and softmax confidence. Incorrect high-confidence predictions are particularly important because they reveal that probability magnitude is not automatically calibrated uncertainty.")
    add_figure(doc, M6 / "evidence" / "04_test_predictions.png", 4, "Representative Held-Out Predictions and Confidence", 6.0)

    doc.add_heading("Limitations and Improvements", level=1)
    add_body(doc, "The model demonstrates a valid end-to-end CNN workflow but should not be treated as production-ready. Accuracy could be improved with longer training after a broader learning-rate search, residual blocks, stronger augmentation, weight decay, or transfer learning from a larger image corpus. Repeated seeded runs would quantify training variance. Calibration metrics and reliability diagrams would help determine whether softmax confidence corresponds to observed correctness. Error analysis should also examine mislabeled or inherently ambiguous images rather than assuming every disagreement represents model failure.")

    doc.add_heading("Conclusion", level=1)
    add_body(doc, f"The CNN learned useful spatial features from the complete CIFAR-10 training partition and achieved {results['test_accuracy'] * 100:.2f}% accuracy on 10,000 held-out images. The validation history, confusion matrix, class-level metrics, and sample predictions provide complementary evidence of model veracity. The result is reproducible through a fixed seed, recorded TensorFlow version, saved metrics, preserved training history, and an exported Keras model. Future work should focus on systematic tuning, calibration, and architecture comparisons rather than relying only on additional epochs.")
    add_references(doc, [
        [("Krizhevsky, A., Nair, V., & Hinton, G. (2009). ", False), ("The CIFAR-10 dataset", True), (" [Data set]. University of Toronto. https://www.cs.toronto.edu/~kriz/cifar.html", False)],
        [("TensorFlow. (2026). ", False), ("Convolutional neural network (CNN)", True), (". https://www.tensorflow.org/tutorials/images/cnn", False)],
    ])
    output = M6 / "CSC580_CTA6_Option_1_Dunn_Justan.docx"
    doc.save(output)
    return output


if __name__ == "__main__":
    print(build_module5())
    print(build_module6())
