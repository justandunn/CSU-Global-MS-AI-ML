"""Build the APA 7 Word report for CSC580 Final Portfolio Option 2."""

from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "CSC580_FinalPortfolio_Option_2_Dunn_Justan.docx"
RESULTS = json.loads((ROOT / "evidence" / "results.json").read_text(encoding="utf-8"))


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    fld_char_1 = OxmlElement("w:fldChar")
    fld_char_1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char_2 = OxmlElement("w:fldChar")
    fld_char_2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char_1, instr_text, fld_char_2])


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    add_page_number(section.header.paragraphs[0])

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 2
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.first_line_indent = Inches(0.5)

    for style_name in ["Title", "Heading 1", "Heading 2", "Heading 3"]:
        style = styles[style_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = True

    styles["Title"].font.size = Pt(14)
    title_ppr = styles["Title"]._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)
    styles["Heading 1"].font.size = Pt(12)
    styles["Heading 1"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    styles["Heading 1"].paragraph_format.space_before = Pt(0)
    styles["Heading 1"].paragraph_format.space_after = Pt(0)
    styles["Heading 2"].font.size = Pt(12)
    styles["Heading 2"].paragraph_format.space_before = Pt(12)
    styles["Heading 2"].paragraph_format.space_after = Pt(0)


def add_body_paragraph(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(text)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT


def add_title_page(document: Document) -> None:
    for _ in range(4):
        document.add_paragraph()
    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Option 2 Encoder Decoder Models in Industry and Sequence Prediction Implementation")
    document.add_paragraph()
    for line in [
        "Justan Dunn",
        "Colorado State University Global",
        "CSC580 Applying Machine Learning and Neural Networks",
        "Dong Nguyen",
        "September 13, 2026",
    ]:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.first_line_indent = Inches(0)
        paragraph.add_run(line)
    document.add_page_break()


def add_research_section(document: Document) -> None:
    document.add_heading(
        "Option 2 Encoder Decoder Models in Industry and Sequence Prediction Implementation",
        level=1,
    )
    add_body_paragraph(
        document,
        "Encoder-decoder neural networks learn a mapping from one structured input to another structured output. "
        "The encoder compresses or contextualizes the source, while the decoder generates the target one element "
        "at a time. This organization is valuable when input and output lengths differ or when each output depends "
        "on both the source and earlier outputs. Across language, speech, visual accessibility, and document services, "
        "encoder-decoder models replace brittle task-specific pipelines with trainable mappings while retaining "
        "important limitations involving data quality, evaluation, and explainability.",
    )

    document.add_heading("Machine Translation and Localization", level=2)
    add_body_paragraph(
        document,
        "Machine translation is the canonical sequence-to-sequence application. An encoder represents a sentence "
        "in one language, and a decoder generates its meaning in another language. Sutskever et al. (2014) showed "
        "that multilayer LSTMs could translate English to French end to end without requiring the fixed phrase tables "
        "used by traditional statistical systems. This capability benefits global commerce, customer support, travel, "
        "and software localization because one architecture can learn reordering and dependencies that span an entire "
        "sentence. Human review remains important for legal, safety, or culturally sensitive content, and evaluation "
        "should include meaning preservation rather than relying only on automated similarity scores.",
    )

    document.add_heading("Conversational Speech Recognition", level=2)
    add_body_paragraph(
        document,
        "Speech recognition maps a variable-length acoustic sequence to a shorter character or word sequence. In the "
        "Listen, Attend and Spell architecture, a pyramidal recurrent encoder processes acoustic filter-bank features, "
        "while an attention-based decoder emits characters (Chan et al., 2016). This end-to-end design reduces the need "
        "for separately engineered acoustic, pronunciation, and language-model components. The business benefit appears "
        "in call-center transcription, meeting notes, voice interfaces, and accessibility tools. However, organizations "
        "must test performance across accents, noise conditions, microphones, and specialized vocabulary because an "
        "average word-error rate can conceal unequal service quality.",
    )

    document.add_heading("Image Captioning and Visual Accessibility", level=2)
    add_body_paragraph(
        document,
        "Encoder-decoder models can also connect different data types. Vinyals et al. (2015) used a convolutional neural "
        "network to encode an image and an LSTM decoder to generate a natural-language description. This pattern can "
        "support alternative text for people with visual impairments, searchable media catalogs, product-description "
        "drafting, and inspection-report assistance. Its benefit over simple image classification is that the output can "
        "describe several objects and relationships instead of selecting one category. Because a fluent caption may still "
        "be factually wrong, deployments should retain confidence thresholds, human correction, and evaluations that "
        "penalize invented objects or missing safety-relevant details.",
    )

    document.add_heading("Abstractive Summarization", level=2)
    add_body_paragraph(
        document,
        "Abstractive summarization encodes a source document and decodes a shorter statement of its meaning. Unlike an "
        "extractive system that copies sentences, an abstractive model can paraphrase, reorder information, and combine "
        "related ideas. Rush et al. (2015) demonstrated an attention-based model that generated each summary word while "
        "conditioning on the source. Summarization can reduce review time in news, research, legal discovery, insurance, "
        "and customer-service operations. The same flexibility also creates risk: a decoder can produce a plausible claim "
        "that the source never made. Effective use therefore requires traceability to source passages, factuality testing, "
        "and human approval when a summary informs consequential decisions.",
    )

    document.add_heading("Research Conclusion", level=2)
    add_body_paragraph(
        document,
        "These applications share a common advantage: the encoder-decoder pattern learns a conditional transformation "
        "between structures instead of forcing inputs and outputs into identical shapes. LSTMs provide memory for ordered "
        "dependencies, and attention can reduce the bottleneck created by one fixed context vector. Nevertheless, industry "
        "value depends on representative data, task-appropriate metrics, monitoring, and human oversight. The implementation "
        "that follows isolates the architecture's core behavior with an exact, auditable sequence-reversal problem.",
    )
    document.add_page_break()


def add_figure(document: Document, path: Path, number: int, title: str, width: float) -> None:
    label = document.add_paragraph()
    label.paragraph_format.first_line_indent = Inches(0)
    label.paragraph_format.keep_with_next = True
    run = label.add_run(f"Figure {number}")
    run.bold = True
    caption = document.add_paragraph()
    caption.paragraph_format.first_line_indent = Inches(0)
    caption.paragraph_format.keep_with_next = True
    title_run = caption.add_run(title)
    title_run.italic = True
    document.add_picture(str(path), width=Inches(width))
    picture_paragraph = document.paragraphs[-1]
    picture_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture_paragraph.paragraph_format.first_line_indent = Inches(0)
    picture_paragraph.paragraph_format.keep_together = True


def add_implementation_section(document: Document) -> None:
    document.add_heading("Programming Implementation", level=1)
    add_body_paragraph(
        document,
        "The program implements the required scalable sequence problem with modern Keras. Each source contains six "
        "integers from 1 through 50. The target contains the first three source integers in reverse order, and zero is "
        "reserved as the decoder start token. Separate random-number generators create 20,000 training, 2,000 validation, "
        "and 100 held-out test sequences. This separation prevents test examples from influencing training and makes the "
        "final exact-match result an independent check.",
    )
    document.add_heading("Processing Design", level=2)
    add_body_paragraph(
        document,
        "Figure 1 documents the complete workflow. Source, shifted decoder-input, and target sequences are one-hot encoded "
        "into 51-category vectors. During teacher-forced training, the decoder receives zero followed by the previous true "
        "target values. The encoder's final hidden and cell states initialize the decoder. During inference, the encoder runs "
        "once; the decoder then predicts one token at a time and feeds that prediction back to generate the three-value target.",
    )
    add_figure(
        document,
        ROOT / "flowchart" / "encoder_decoder_lstm_flowchart.png",
        1,
        "Encoder Decoder LSTM Processing Workflow",
        5.45,
    )

    document.add_page_break()
    document.add_heading("Model Configuration and Training", level=2)
    config = RESULTS["configuration"]
    table = document.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    headers = table.rows[0].cells
    headers[0].text = "Configuration Item"
    headers[1].text = "Value"
    for cell in headers:
        set_cell_shading(cell, "D9EAF7")
        for run in cell.paragraphs[0].runs:
            run.bold = True
    for label, value in [
        ("Input and output vocabulary", "51 tokens including start token 0"),
        ("Sequence dimensions", "6 input steps and 3 output steps"),
        ("Encoder and decoder capacity", f"{config['lstm_units']} LSTM units each"),
        ("Training and validation samples", f"{config['training_samples']:,} and {config['validation_samples']:,}"),
        ("Optimizer and loss", "Adam and categorical cross-entropy"),
        ("Batch size", str(config["batch_size"])),
        ("Epochs completed", str(config["epochs_completed"])),
        ("Random seed", str(RESULTS["seed"])),
    ]:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = value
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.first_line_indent = Inches(0)
                paragraph.paragraph_format.line_spacing = 1
                paragraph.paragraph_format.space_after = Pt(2)
                for run in paragraph.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(10)

    add_body_paragraph(
        document,
        "The model used 128 LSTM units in both recurrent components and a 51-unit softmax layer for token probabilities. "
        "Adam optimized categorical cross-entropy. Early stopping monitored validation loss with a patience of five epochs "
        "and restored the best weights. The loss continued to improve through the configured 40th epoch, so all epochs were "
        "completed. Figure 2 shows smooth convergence without a material divergence between training and validation curves.",
    )
    add_figure(document, ROOT / "evidence" / "02_training_history.png", 2, "Training and Validation History", 6.2)

    document.add_page_break()
    document.add_heading("Evaluation Results", level=2)
    evaluation = RESULTS["evaluation"]
    training = RESULTS["training"]
    add_body_paragraph(
        document,
        f"The final validation token accuracy was {training['final_validation_accuracy']:.2%}, with validation loss "
        f"{training['final_validation_loss']:.6f}. Autoregressive inference then evaluated 100 newly generated sequences. "
        f"The model predicted {evaluation['correct_sequences']} entire targets correctly, producing "
        f"{evaluation['exact_match_accuracy']:.2%} exact-match accuracy. Exact matching is stricter than token accuracy: "
        "all three values must appear in the correct order for a sequence to count as correct. Figure 3 provides ten "
        "representative held-out examples, while the accompanying runtime text file preserves all 100 predictions.",
    )
    add_figure(document, ROOT / "evidence" / "03_prediction_examples.png", 3, "Representative Held Out Predictions", 6.25)
    add_body_paragraph(
        document,
        "Ninety-nine percent exact accuracy exceeds the project's 95% acceptance threshold and demonstrates that the "
        "encoder states retained the relevant source order while the inference decoder generated the reversed target. "
        "The single error occurred for source [43, 42, 3, 42, 2, 7]: the expected target was [3, 42, 43], but the model "
        "returned [42, 3, 43]. This was not a data-pipeline error; it was a model prediction error involving a repeated "
        "value later in the source. Additional training examples, attention, or beam search could improve this edge case.",
    )

    document.add_heading("Reproducibility and Limitations", level=2)
    add_body_paragraph(
        document,
        "The source records the Python, TensorFlow, Keras, and NumPy versions; fixed seed; tensor shapes; hyperparameters; "
        "and exact evaluation procedure. The generated results JSON and prediction text provide machine-readable and "
        "human-readable evidence. The principal limitation is that the task is synthetic and substantially simpler than "
        "translation, speech recognition, or summarization. Its controlled target is useful for validating architecture "
        "and inference logic, but it does not establish performance on ambiguous natural data. A production extension "
        "should add attention, masks for variable-length inputs, beam-search decoding, domain-specific metrics, and drift monitoring.",
    )

    document.add_heading("Conclusion", level=2)
    add_body_paragraph(
        document,
        "The project satisfies both required components of Option 2. The research demonstrates four distinct industry "
        "uses for encoder-decoder models, and the programming implementation trains and independently evaluates a documented "
        "Keras LSTM. The 99% exact-match result confirms that the model learned the required transformation while the one "
        "observed error provides a defensible basis for further improvement. The flowchart, source code, runtime output, "
        "figures, and versioned metrics make the submission reproducible and auditable.",
    )


def add_references(document: Document) -> None:
    document.add_page_break()
    document.add_heading("References", level=1)
    references = [
        "Chan, W., Jaitly, N., Le, Q. V., & Vinyals, O. (2016). Listen, attend and spell: A neural network for large vocabulary conversational speech recognition. In 2016 IEEE International Conference on Acoustics, Speech and Signal Processing (ICASSP) (pp. 4960-4964). IEEE. https://doi.org/10.1109/ICASSP.2016.7472621",
        "Rush, A. M., Chopra, S., & Weston, J. (2015). A neural attention model for abstractive sentence summarization. In Proceedings of the 2015 Conference on Empirical Methods in Natural Language Processing (pp. 379-389). Association for Computational Linguistics. https://doi.org/10.18653/v1/D15-1044",
        "Sutskever, I., Vinyals, O., & Le, Q. V. (2014). Sequence to sequence learning with neural networks. Advances in Neural Information Processing Systems, 27, 3104-3112. https://papers.nips.cc/paper/5346-sequence-to-sequence-learning-with-neural-networks",
        "Vinyals, O., Toshev, A., Bengio, S., & Erhan, D. (2015). Show and tell: A neural image caption generator. In Proceedings of the IEEE Conference on Computer Vision and Pattern Recognition (pp. 3156-3164). https://doi.org/10.1109/CVPR.2015.7298935",
    ]
    for reference in references:
        paragraph = document.add_paragraph(reference)
        paragraph.paragraph_format.first_line_indent = Inches(-0.5)
        paragraph.paragraph_format.left_indent = Inches(0.5)
        paragraph.paragraph_format.line_spacing = 2
        paragraph.paragraph_format.space_after = Pt(0)


def main() -> None:
    document = Document()
    configure_document(document)
    add_title_page(document)
    add_research_section(document)
    add_implementation_section(document)
    add_references(document)
    document.core_properties.title = (
        "Option 2 Encoder Decoder Models in Industry and Sequence Prediction Implementation"
    )
    document.core_properties.author = "Justan Dunn"
    document.core_properties.subject = "CSC580 Final Portfolio Project Option 2"
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
