"""Build APA-style Word analyses for both CSC580 Module 4 Option 2 submissions."""
from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parent
PM = ROOT / "portfolio_milestone_option_2"
CT = ROOT / "critical_thinking_option_2"


def font(run, name="Times New Roman", size=12, bold=None, italic=None):
    run.font.name = name; run.font.size = Pt(size)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if bold is not None: run.bold = bold
    if italic is not None: run.italic = italic


def page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run(); font(run)
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    text = OxmlElement("w:instrText"); text.set(qn("xml:space"), "preserve"); text.text = " PAGE "
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, text, end])


def setup(title, subject):
    doc = Document(); sec = doc.sections[0]
    for attr in ("top_margin", "bottom_margin", "left_margin", "right_margin"): setattr(sec, attr, Inches(1))
    page_number(sec.header.paragraphs[0])
    normal = doc.styles["Normal"]; normal.font.name = "Times New Roman"; normal.font.size = Pt(12)
    normal._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal.paragraph_format.line_spacing = 2; normal.paragraph_format.space_after = Pt(0)
    for name in ("Heading 1", "Heading 2"):
        s = doc.styles[name]; s.font.name = "Times New Roman"; s.font.size = Pt(12); s.font.bold = True
        s.font.color.rgb = None; s._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Times New Roman")
        s._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Times New Roman")
    doc.styles["Heading 1"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for _ in range(4): doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; font(p.add_run(title), bold=True)
    for line in ("Justan Dunn", "Colorado State University Global",
                 "CSC580-1: Applying Machine Learning and Neural Networks", "Dong Nguyen", "August 17, 2026"):
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; font(p.add_run(line))
    doc.add_page_break(); p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; font(p.add_run(title), bold=True)
    doc.core_properties.title = title; doc.core_properties.author = "Justan Dunn"; doc.core_properties.subject = subject
    return doc


def body(doc, text):
    p = doc.add_paragraph(text); p.paragraph_format.first_line_indent = Inches(.5)


def figure(doc, folder, filename, number, title, width=6.1):
    p = doc.add_paragraph(); p.paragraph_format.keep_with_next = True; p.paragraph_format.line_spacing = 1
    font(p.add_run(f"Figure {number}\n"), bold=True); font(p.add_run(title), italic=True)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(6)
    p.add_run().add_picture(str(folder / filename), width=Inches(width))


def references(doc, entries):
    doc.add_page_break(); doc.add_heading("References", 1)
    for parts in entries:
        p = doc.add_paragraph(); p.paragraph_format.left_indent = Inches(.5); p.paragraph_format.first_line_indent = Inches(-.5)
        for text, italic in parts: font(p.add_run(text), italic=italic)


def build_pm():
    r = json.loads((PM / "evidence" / "results.json").read_text())
    title = "Improving TensorFlow Model Performance and Quality"
    doc = setup(title, "CSC580 Module 4 Portfolio Milestone, Option 2")
    doc.add_heading("Purpose and Method", 1)
    body(doc, "This portfolio milestone improved the Module 3 Auto MPG neural-network regression model by preventing unnecessary training and evaluating generalization on data excluded from model development. The workflow retained the cleaned 392-row Auto MPG dataset, deterministic 80/20 split, one-hot encoding of vehicle origin, and normalization based only on training statistics (UCI Machine Learning Repository, 1993). The Keras network contained two 64-neuron rectified linear hidden layers and one continuous MPG output. An EarlyStopping callback monitored validation loss with a patience of 10 epochs and restored the weights from the best epoch (TensorFlow, 2026).")
    doc.add_heading("Early-Stopping Results", 1)
    body(doc, f"Training stopped automatically after {r['epochs_completed']} epochs instead of consuming the 1,000-epoch limit. The best validation loss occurred at epoch {r['best_epoch']}, and the best validation MAE was {r['best_validation_mae']:.2f} MPG. Figure 1 shows rapid early learning followed by a stable validation curve. The approximately 2.30-MPG validation error is useful for estimating broad fuel-efficiency differences, although it is not precise enough to treat every individual prediction as exact. Early stopping addresses the degradation observed in the earlier long run and reduces both computation and overfitting risk.")
    figure(doc, PM / "evidence", "01_early_stopping_mae.png", 1, "Training and Validation MAE With Early Stopping")
    doc.add_heading("Held-Out Test Evaluation", 1)
    body(doc, f"The untouched {r['test_rows']}-vehicle test set produced an MAE of {r['test_mae']:.2f} MPG, MSE of {r['test_mse']:.2f} MPG squared, and RMSE of {r['test_rmse']:.2f} MPG. MAE means that a typical prediction differed from the observed value by about 1.79 MPG. This is slightly better than validation performance and indicates that the selected weights generalized without obvious collapse. RMSE is higher because squaring gives additional influence to the relatively small number of large misses. The test result is credible because neither test features nor labels influenced normalization, training, early stopping, or model selection.")
    figure(doc, PM / "evidence", "02_test_evaluation.png", 2, "Held-Out Test-Set Evaluation Output", 6.0)
    doc.add_heading("Prediction Quality", 1)
    body(doc, "Figure 3 compares every predicted MPG value with its true test label. Most observations cluster near the diagonal reference line, showing that the network learned the principal relationship between vehicle specifications and fuel economy. Dispersion increases among some higher-MPG vehicles, where the historical sample is thinner and vehicle designs may not be fully represented by the retained predictors. The model is therefore appropriate as a screening or estimation tool, but unusually efficient vehicles warrant a confidence interval or human review.")
    figure(doc, PM / "evidence", "03_true_vs_predicted.png", 3, "True MPG Compared With Model Predictions", 5.5)
    doc.add_heading("Error Distribution", 1)
    body(doc, f"Prediction error was defined as predicted MPG minus true MPG. The mean error was {r['mean_error']:.2f} MPG and the median was {r['median_error']:.2f} MPG, both close to zero; consequently, the model showed little overall directional bias. The error standard deviation was {r['error_std']:.2f} MPG. Figure 4 is approximately mound-shaped around zero but is not perfectly normal because its tails and individual extreme errors are asymmetric. Normality is a useful descriptive check, not a requirement for neural-network regression. The near-zero center is encouraging, while the tails reinforce the need to report MAE and RMSE together.")
    figure(doc, PM / "evidence", "04_error_distribution.png", 4, "Distribution of Test-Set Prediction Errors", 6.0)
    doc.add_heading("Conclusion", 1)
    body(doc, "The milestone successfully corrected the principal weakness identified in Module 3. Validation-based early stopping selected a substantially earlier model, restored its best weights, and achieved a held-out MAE of 1.79 MPG. The prediction and residual plots show useful generalization with limited aggregate bias, although isolated large errors remain. Future work should repeat the experiment across several fixed splits, quantify uncertainty, and examine the largest residuals by vehicle origin and model year before operational deployment.")
    references(doc, [[("TensorFlow. (2026). ",False),("tf.keras.callbacks.EarlyStopping",True),(". https://www.tensorflow.org/api_docs/python/tf/keras/callbacks/EarlyStopping",False)],[("UCI Machine Learning Repository. (1993). ",False),("Auto MPG",True),(" [Data set]. https://doi.org/10.24432/C5859H",False)]])
    path = PM / "CSC580_MidTermPortfolio_Option_2_Dunn_Justan.docx"; doc.save(path); return path


def build_ct():
    r = json.loads((CT / "evidence" / "results.json").read_text())
    title = "Logistic Regression With TensorFlow"
    doc = setup(title, "CSC580 Module 4 Critical Thinking Assignment, Option 2")
    doc.add_heading("Purpose and Data Generation", 1)
    body(doc, "This project implemented binary logistic regression in TensorFlow using 100 reproducible synthetic observations. Fifty class-zero points were sampled from a two-dimensional Gaussian distribution centered at (-1, -1), and fifty class-one points were sampled from a Gaussian distribution centered at (1, 1). Both distributions used identity covariance, so overlap was expected near the center. Figure 1 confirms that the classes are generally separated while retaining several ambiguous observations suitable for evaluating probabilistic classification.")
    figure(doc, CT / "evidence", "01_input_classes.png", 1, "Synthetic Gaussian Class-Zero and Class-One Observations", 5.6)
    doc.add_heading("TensorFlow Model", 1)
    body(doc, f"The TensorFlow {r['tensorflow_version']} graph accepted two-feature inputs and connected them to one sigmoid output. The sigmoid transformed the linear combination of inputs, weights, and bias into probabilities between zero and one. Binary cross-entropy quantified disagreement between probabilities and labels, and Adam optimization updated the parameters for 200 epochs. A TensorBoard callback wrote the graph and epoch summaries to the included logs directory. The learned weights were {r['weights'][0]:.3f} and {r['weights'][1]:.3f}, with a bias of {r['bias']:.3f}; both positive weights are consistent with class one occupying the upper-right region (TensorFlow, 2026).")
    doc.add_heading("Predictions and Decision Boundary", 1)
    body(doc, f"Probabilities of at least 0.50 were assigned to class one. The trained model achieved {r['accuracy']*100:.1f}% accuracy and cross-entropy loss of {r['loss']:.3f}, misclassifying {r['misclassified']} of {r['samples']} observations. Figure 2 overlays predicted classes and the 0.50 decision boundary on the data. Most points fall on the expected side of the line. The remaining mistakes occur primarily in the overlap region and represent reasonable uncertainty rather than a failure to learn the dominant pattern.")
    figure(doc, CT / "evidence", "02_predictions_boundary.png", 2, "Predicted Classes and the Learned Probability Boundary", 5.7)
    doc.add_heading("Training Behavior and Interpretation", 1)
    body(doc, "Figure 3 shows that cross-entropy declined quickly and then stabilized while accuracy approached its final level. This pattern indicates convergence rather than continued oscillation or divergence. Because the two generating distributions overlap, forcing perfect training accuracy would be undesirable: a much more flexible classifier could memorize noise without improving generalization. Logistic regression instead provides an interpretable linear boundary and calibrated probability structure appropriate for this synthetic problem.")
    figure(doc, CT / "evidence", "03_training_history.png", 3, "Cross-Entropy Loss and Accuracy Across Training Epochs", 6.0)
    doc.add_heading("Conclusion", 1)
    body(doc, "The TensorFlow logistic-regression model correctly represented the assignment's binary classification workflow from data creation through training, prediction, visualization, and summary logging. Its 94.0% accuracy, low cross-entropy, and correctly oriented boundary demonstrate strong performance while acknowledging unavoidable Gaussian overlap. The implementation also remains reproducible through fixed seeds, recorded package version, saved evidence, and TensorBoard event logs.")
    references(doc, [[("TensorFlow. (2026). ",False),("tf.keras API",True),(". https://www.tensorflow.org/api_docs/python/tf/keras",False)]])
    path = CT / "CSC580_CTA_4_2_Dunn_Justan.docx"; doc.save(path); return path


if __name__ == "__main__":
    print(build_pm()); print(build_ct())
