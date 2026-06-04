from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer


OUTPUT_DIR = Path(__file__).resolve().parent
DOCX_PATH = OUTPUT_DIR / "CSC525_Module3PM_Dunn_Justan.docx"
PDF_PATH = OUTPUT_DIR / "CSC525_Module3PM_Dunn_Justan.pdf"


TITLE = "Option #2: NLP Chatbot Project Draft"
RUNNING_TITLE = "NLP Chatbot Project Draft"


BODY_PARAGRAPHS = [
    (
        "For the final portfolio project, the selected option is Option #2, the NLP Chatbot Project. "
        "The proposed system is a closed-domain enterprise knowledge-support chatbot designed to help "
        "employees retrieve accurate answers from a curated company knowledge base. The initial use case "
        "will focus on common operational questions such as policy lookup, process guidance, onboarding "
        "support, and frequently requested internal procedures. This direction is intentionally practical: "
        "many organizations have information stored across manuals, standard operating procedures, training "
        "documents, and frequently asked question lists, but employees often struggle to locate the right "
        "answer quickly. The chatbot's goal will be to interpret a user's natural-language question, identify "
        "the most relevant knowledge-base entry, and return a clear response with enough context for the "
        "employee to act."
    ),
    (
        "The chatbot will demonstrate adaptive control in an uncertain environment by treating each user "
        "question as an uncertain input condition. The system will not assume that users know the exact policy "
        "name, document title, or vocabulary used in the knowledge base. Instead, it will adapt its response "
        "strategy based on the similarity between the user's question and the available training entries. "
        "When confidence is high, the bot will return the best-matching answer. When confidence is moderate, "
        "it will return a cautious response and may suggest the closest matching topic. When confidence is "
        "low, it will ask the user to rephrase or narrow the question rather than inventing an unsupported "
        "answer. This confidence-based response control is important because a business chatbot should avoid "
        "nonsense responses and should make uncertainty visible to the user."
    ),
    (
        "The first version will be a retrieval-based chatbot rather than an open-domain generative chatbot. "
        "The planned implementation will use Python, pandas for loading the knowledge-base dataset, and "
        "scikit-learn for text vectorization and similarity scoring. The core model will use term frequency-"
        "inverse document frequency (TF-IDF) features and cosine similarity to compare the user's input with "
        "stored questions and answers. TF-IDF is a reasonable first model because it is explainable, fast, "
        "and well suited for matching user queries against a controlled document collection (Manning et al., "
        "2008). Scikit-learn is also appropriate for this coursework prototype because it provides a stable "
        "Python machine learning interface and includes text feature extraction tools such as TfidfVectorizer "
        "(Pedregosa et al., 2011; scikit-learn developers, 2025)."
    ),
    (
        "The program structure will include a data folder with a CSV knowledge base, a preprocessing module, "
        "a retrieval model module, a chatbot interface, and a small testing script. The dataset will contain "
        "fields such as topic, sample question, approved answer, source category, and optional keywords. "
        "The preprocessing step will normalize case, remove unnecessary punctuation, and prepare text for "
        "vectorization. The retrieval model will fit the TF-IDF vectorizer on the knowledge-base entries and "
        "rank candidate answers using cosine similarity. The chatbot interface will accept user input from "
        "the command line or a lightweight web interface, display the selected answer, and show a confidence "
        "score for testing. Initial hyperparameters will include unigram and bigram features, a maximum "
        "document-frequency threshold to reduce overly common terms, and a minimum similarity threshold for "
        "fallback behavior."
    ),
    (
        "The final version may compare this baseline with a sentence-embedding approach or transformer-based "
        "model if time and dependency setup permit. Transformer models are important in modern NLP because "
        "attention-based architectures improved the ability of models to represent language context (Vaswani "
        "et al., 2017). However, for the final course deliverable, the safest design is a closed-domain "
        "retrieval system that can be tested with known expected answers. The project will be evaluated by "
        "running sample employee questions through the chatbot, recording whether the retrieved answers are "
        "relevant, and identifying where additional training examples or better threshold tuning are needed. "
        "This approach creates a defensible prototype for later enterprise use."
    ),
]


REFERENCES = [
    "Manning, C. D., Raghavan, P., & Schutze, H. (2008). Introduction to information retrieval. Cambridge University Press. https://nlp.stanford.edu/IR-book/",
    "Pedregosa, F., Varoquaux, G., Gramfort, A., Michel, V., Thirion, B., Grisel, O., Blondel, M., Prettenhofer, P., Weiss, R., Dubourg, V., Vanderplas, J., Passos, A., Cournapeau, D., Brucher, M., Perrot, M., & Duchesnay, E. (2011). Scikit-learn: Machine learning in Python. Journal of Machine Learning Research, 12, 2825-2830. https://jmlr.csail.mit.edu/papers/v12/pedregosa11a.html",
    "scikit-learn developers. (2025). TfidfVectorizer. scikit-learn 1.8.0 documentation. https://scikit-learn.org/stable/modules/generated/sklearn.feature_extraction.text.TfidfVectorizer.html",
    "Vaswani, A., Shazeer, N., Parmar, N., Uszkoreit, J., Jones, L., Gomez, A. N., Kaiser, L., & Polosukhin, I. (2017). Attention is all you need. Advances in Neural Information Processing Systems, 30. https://arxiv.org/abs/1706.03762",
]


def add_page_number(paragraph) -> None:
    """Add a PAGE field to a paragraph for APA-style page numbering."""
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()

    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")

    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"

    fld_char_separate = OxmlElement("w:fldChar")
    fld_char_separate.set(qn("w:fldCharType"), "separate")

    text = OxmlElement("w:t")
    text.text = "1"

    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")

    run._r.append(fld_char_begin)
    run._r.append(instr_text)
    run._r.append(fld_char_separate)
    run._r.append(text)
    run._r.append(fld_char_end)


def set_document_defaults(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.5)
    section.footer_distance = Inches(0.5)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 2
    normal.paragraph_format.space_after = Pt(0)

    for style_name in ["Title", "Heading 1", "Heading 2"]:
        style = doc.styles[style_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.color.rgb = None


def add_centered_line(doc: Document, text: str, bold: bool = False) -> None:
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.line_spacing = 2
    run = para.add_run(text)
    run.bold = bold
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(12)


def add_body_paragraph(doc: Document, text: str) -> None:
    para = doc.add_paragraph()
    para.paragraph_format.first_line_indent = Inches(0.5)
    para.paragraph_format.line_spacing = 2
    para.paragraph_format.space_after = Pt(0)
    run = para.add_run(text)
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(12)


def add_section_heading(doc: Document, text: str) -> None:
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    para.paragraph_format.line_spacing = 2
    para.paragraph_format.space_after = Pt(0)
    run = para.add_run(text)
    run.bold = True
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(12)


def add_reference(doc: Document, text: str) -> None:
    para = doc.add_paragraph()
    para.paragraph_format.left_indent = Inches(0.5)
    para.paragraph_format.first_line_indent = Inches(-0.5)
    para.paragraph_format.line_spacing = 2
    para.paragraph_format.space_after = Pt(0)
    run = para.add_run(text)
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(12)


def build_document() -> None:
    doc = Document()
    set_document_defaults(doc)

    header_para = doc.sections[0].header.paragraphs[0]
    add_page_number(header_para)

    for _ in range(4):
        doc.add_paragraph()

    add_centered_line(doc, TITLE, bold=True)
    add_centered_line(doc, RUNNING_TITLE)
    add_centered_line(doc, "Justan Dunn")
    add_centered_line(doc, "Colorado State University Global")
    add_centered_line(doc, "CSC525: Principles of Machine Learning")
    add_centered_line(doc, "Module 3 Portfolio Milestone")
    add_centered_line(doc, "May 31, 2026")

    doc.add_page_break()

    title_para = doc.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_para.paragraph_format.line_spacing = 2
    title_run = title_para.add_run(TITLE)
    title_run.bold = True
    title_run.font.name = "Times New Roman"
    title_run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    title_run.font.size = Pt(12)

    add_body_paragraph(doc, BODY_PARAGRAPHS[0])
    add_body_paragraph(doc, BODY_PARAGRAPHS[1])

    add_section_heading(doc, "Technical Plan")
    for paragraph in BODY_PARAGRAPHS[2:]:
        add_body_paragraph(doc, paragraph)

    doc.add_page_break()

    refs_heading = doc.add_paragraph()
    refs_heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    refs_heading.paragraph_format.line_spacing = 2
    refs_run = refs_heading.add_run("References")
    refs_run.bold = True
    refs_run.font.name = "Times New Roman"
    refs_run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    refs_run.font.size = Pt(12)

    for reference in REFERENCES:
        add_reference(doc, reference)

    doc.save(DOCX_PATH)
    print(DOCX_PATH)


def on_page(canvas, _doc) -> None:
    canvas.saveState()
    canvas.setFont("Times-Roman", 12)
    canvas.drawRightString(7.5 * inch, 10.5 * inch, str(canvas.getPageNumber()))
    canvas.restoreState()


def build_pdf() -> None:
    doc = SimpleDocTemplate(
        str(PDF_PATH),
        pagesize=letter,
        rightMargin=inch,
        leftMargin=inch,
        topMargin=inch,
        bottomMargin=inch,
    )

    title_style = ParagraphStyle(
        "TitleStyle",
        fontName="Times-Bold",
        fontSize=12,
        leading=24,
        alignment=TA_CENTER,
        spaceAfter=0,
    )
    center_style = ParagraphStyle(
        "CenterStyle",
        fontName="Times-Roman",
        fontSize=12,
        leading=24,
        alignment=TA_CENTER,
        spaceAfter=0,
    )
    body_style = ParagraphStyle(
        "BodyStyle",
        fontName="Times-Roman",
        fontSize=12,
        leading=24,
        firstLineIndent=0.5 * inch,
        alignment=TA_LEFT,
        spaceAfter=0,
    )
    heading_style = ParagraphStyle(
        "HeadingStyle",
        fontName="Times-Bold",
        fontSize=12,
        leading=24,
        alignment=TA_LEFT,
        spaceBefore=0,
        spaceAfter=0,
    )
    refs_style = ParagraphStyle(
        "ReferencesStyle",
        fontName="Times-Roman",
        fontSize=12,
        leading=24,
        leftIndent=0.5 * inch,
        firstLineIndent=-0.5 * inch,
        alignment=TA_LEFT,
        spaceAfter=0,
    )

    story = []
    story.append(Spacer(1, 1.5 * inch))
    story.append(Paragraph(TITLE, title_style))
    story.append(Paragraph(RUNNING_TITLE, center_style))
    story.append(Paragraph("Justan Dunn", center_style))
    story.append(Paragraph("Colorado State University Global", center_style))
    story.append(Paragraph("CSC525: Principles of Machine Learning", center_style))
    story.append(Paragraph("Module 3 Portfolio Milestone", center_style))
    story.append(Paragraph("May 31, 2026", center_style))
    story.append(PageBreak())

    story.append(Paragraph(TITLE, title_style))
    story.append(Paragraph(BODY_PARAGRAPHS[0], body_style))
    story.append(Paragraph(BODY_PARAGRAPHS[1], body_style))
    story.append(Paragraph("Technical Plan", heading_style))
    for paragraph in BODY_PARAGRAPHS[2:]:
        story.append(Paragraph(paragraph, body_style))

    story.append(PageBreak())
    story.append(Paragraph("References", title_style))
    for reference in REFERENCES:
        story.append(Paragraph(reference, refs_style))

    doc.build(story, onFirstPage=on_page, onLaterPages=on_page)
    print(PDF_PATH)


if __name__ == "__main__":
    build_document()
    build_pdf()
