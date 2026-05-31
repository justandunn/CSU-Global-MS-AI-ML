"""Build the Module 3 Critical Thinking Project report."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
REPORT_DIR = ROOT / "report"
OUTPUTS_DIR = ROOT / "outputs"
DOCX_PATH = REPORT_DIR / "module3_ctp_report.docx"
CONSOLE_IMAGE = OUTPUTS_DIR / "console_output_screenshot.png"
CODE_IMAGE = OUTPUTS_DIR / "code_excerpt_screenshot.png"
PLOT_IMAGE = OUTPUTS_DIR / "polynomial_regression_fit.png"


def set_cell_text(paragraph, text: str, *, bold: bool = False, italic: bool = False) -> None:
    run = paragraph.add_run(text)
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(12)
    run.bold = bold
    run.italic = italic


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_begin)
    run._r.append(instr_text)
    run._r.append(fld_char_end)


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.5)
    section.footer_distance = Inches(0.5)

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 2.0
    normal.paragraph_format.space_after = Pt(0)

    for style_name in ("Title", "Heading 1", "Heading 2"):
        style = document.styles[style_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.size = Pt(12)
        style.font.bold = True
        style.font.color.rgb = None
        style.paragraph_format.line_spacing = 2.0
        style.paragraph_format.space_after = Pt(0)

    add_page_number(section.header.paragraphs[0])


def add_centered_line(document: Document, text: str, *, bold: bool = False) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.line_spacing = 2.0
    set_cell_text(paragraph, text, bold=bold)


def add_body_paragraph(document: Document, text: str, *, indent: bool = True) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.line_spacing = 2.0
    paragraph.paragraph_format.space_after = Pt(0)
    if indent:
        paragraph.paragraph_format.first_line_indent = Inches(0.5)
    set_cell_text(paragraph, text)


def add_section_heading(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.line_spacing = 2.0
    paragraph.paragraph_format.space_after = Pt(0)
    set_cell_text(paragraph, text, bold=True)


def add_figure(document: Document, figure_number: int, title: str, image_path: Path, note: str, width: float = 6.25) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.line_spacing = 2.0
    set_cell_text(paragraph, f"Figure {figure_number}", bold=True)

    title_paragraph = document.add_paragraph()
    title_paragraph.paragraph_format.line_spacing = 2.0
    set_cell_text(title_paragraph, title, italic=True)

    picture_paragraph = document.add_paragraph()
    picture_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture_paragraph.paragraph_format.line_spacing = 1.0
    picture_paragraph.add_run().add_picture(str(image_path), width=Inches(width))

    note_paragraph = document.add_paragraph()
    note_paragraph.paragraph_format.line_spacing = 2.0
    note_paragraph.paragraph_format.space_after = Pt(0)
    set_cell_text(note_paragraph, "Note. ", italic=True)
    set_cell_text(note_paragraph, note)


def add_reference(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.line_spacing = 2.0
    paragraph.paragraph_format.left_indent = Inches(0.5)
    paragraph.paragraph_format.first_line_indent = Inches(-0.5)
    paragraph.paragraph_format.space_after = Pt(0)
    set_cell_text(paragraph, text)


def build_report() -> None:
    document = Document()
    configure_document(document)

    # APA 7 student-style title page.
    for _ in range(4):
        document.add_paragraph()
    add_centered_line(document, "Simple Polynomial Regression in Python:", bold=True)
    add_centered_line(document, "Predicting Salary From Years of Experience", bold=True)
    document.add_paragraph()
    add_centered_line(document, "Student Name")
    add_centered_line(document, "Colorado State University Global")
    add_centered_line(document, "CSC525: Principles of Machine Learning")
    add_centered_line(document, "Dong Nguyen")
    add_centered_line(document, "May 30, 2026")

    document.add_page_break()

    add_section_heading(document, "Simple Polynomial Regression in Python")
    add_body_paragraph(
        document,
        "This project used polynomial regression to predict salary from years of experience. "
        "Polynomial regression extends a basic linear relationship by adding powered terms for the independent "
        "variable, which allows the fitted model to represent mild curvature in the data (W3Schools, n.d.). "
        "The dataset contained 30 observations with two fields: years of experience and salary (rohankayan, n.d.).",
    )
    add_body_paragraph(
        document,
        "The Python script validates the local CSV file, separates the data into training and test sets, fits a "
        "second-degree polynomial using NumPy least squares, and reports performance on the held-out test set. "
        "NumPy was used for the polynomial fit and numerical evaluation because it provides efficient array "
        "operations for scientific computing (Harris et al., 2020).",
    )

    add_section_heading(document, "Results")
    add_body_paragraph(
        document,
        "The model produced an R-squared value of .9572 on the held-out test data, with a root mean squared error "
        "of $5,026.08 and a mean absolute error of $3,835.38. These results indicate that a simple degree-2 "
        "polynomial model captured most of the relationship between experience and salary in this small dataset. "
        "The output also demonstrates salary predictions for candidates with 2.5, 5.0, 7.5, and 10.0 years of experience.",
    )
    add_figure(
        document,
        1,
        "Console Output From the Polynomial Regression Script",
        CONSOLE_IMAGE,
        "The primary output screenshot shows the model equation, train/test split, evaluation metrics, and sample salary predictions.",
        width=6.25,
    )

    add_body_paragraph(
        document,
        "The fitted curve visually follows the upward relationship in the observed data. However, the model is limited "
        "by the small number of records and by the fact that real salary outcomes depend on many additional variables, "
        "including role, location, education, industry, and market conditions.",
    )
    add_figure(
        document,
        2,
        "Observed Salary Values and Degree-2 Polynomial Fit",
        PLOT_IMAGE,
        "The blue points represent observed salaries, and the red curve represents the polynomial regression fit.",
        width=6.25,
    )

    add_section_heading(document, "Code Evidence")
    add_body_paragraph(
        document,
        "The executable Python file is submitted separately. Figure 3 shows the main model training, evaluation, and "
        "prediction section of the script.",
    )
    add_figure(
        document,
        3,
        "Code Excerpt Showing Model Training and Output Generation",
        CODE_IMAGE,
        "The excerpt shows the main workflow used to load data, fit the polynomial model, evaluate it, and print predictions.",
        width=6.25,
    )

    document.add_page_break()
    add_section_heading(document, "References")
    add_reference(
        document,
        "Harris, C. R., Millman, K. J., van der Walt, S. J., Gommers, R., Virtanen, P., Cournapeau, D., "
        "Wieser, E., Taylor, J., Berg, S., Smith, N. J., Kern, R., Picus, M., Hoyer, S., van Kerkwijk, "
        "M. H., Brett, M., Haldane, A., Fernandez del Rio, J., Wiebe, M., Peterson, P., ... Oliphant, "
        "T. E. (2020). Array programming with NumPy. Nature, 585(7825), 357-362. "
        "https://doi.org/10.1038/s41586-020-2649-2",
    )
    add_reference(
        document,
        "rohankayan. (n.d.). Years of experience and salary dataset [Data set]. Kaggle. "
        "https://www.kaggle.com/datasets/rohankayan/years-of-experience-and-salary-dataset",
    )
    add_reference(
        document,
        "W3Schools. (n.d.). Machine learning - Polynomial regression. "
        "https://www.w3schools.com/python/python_ml_polynomial_regression.asp",
    )

    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    document.save(DOCX_PATH)
    print(DOCX_PATH)


if __name__ == "__main__":
    build_report()
