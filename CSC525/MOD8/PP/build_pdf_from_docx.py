"""Generate a visually reviewable PDF copy of the final APA DOCX report."""

from __future__ import annotations

from html import escape
from pathlib import Path

from docx import Document
from docx.table import Table as DocxTable
from docx.text.paragraph import Paragraph as DocxParagraph
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import (
    KeepTogether,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


BASE_DIR = Path(__file__).resolve().parent
REPORT_DIR = BASE_DIR / "report"
DOCX_PATH = REPORT_DIR / "CSC525_Module8PP_Dunn_Justan.docx"
PDF_PATH = REPORT_DIR / "CSC525_Module8PP_Dunn_Justan.pdf"


BODY = ParagraphStyle(
    "APA Body",
    fontName="Times-Roman",
    fontSize=12,
    leading=24,
    alignment=TA_LEFT,
    firstLineIndent=0.5 * inch,
    spaceBefore=0,
    spaceAfter=0,
    allowWidows=1,
    allowOrphans=1,
)
BODY_NO_INDENT = ParagraphStyle(
    "APA Body No Indent",
    parent=BODY,
    firstLineIndent=0,
)
HEADING_1 = ParagraphStyle(
    "APA Heading 1",
    parent=BODY_NO_INDENT,
    fontName="Times-Bold",
    alignment=TA_CENTER,
    spaceBefore=12,
    keepWithNext=True,
)
HEADING_2 = ParagraphStyle(
    "APA Heading 2",
    parent=HEADING_1,
    alignment=TA_LEFT,
)
TITLE_STYLE = ParagraphStyle(
    "APA Title",
    parent=HEADING_1,
    spaceBefore=0,
    spaceAfter=0,
)
REFERENCE = ParagraphStyle(
    "APA Reference",
    parent=BODY_NO_INDENT,
    leftIndent=0.5 * inch,
    firstLineIndent=-0.5 * inch,
)
CODE = ParagraphStyle(
    "Code",
    fontName="Courier",
    fontSize=10,
    leading=13,
    leftIndent=0.5 * inch,
    rightIndent=0.5 * inch,
    spaceBefore=3,
    spaceAfter=3,
)
TABLE_LABEL = ParagraphStyle(
    "Table Label",
    fontName="Times-Bold",
    fontSize=12,
    leading=14,
    spaceBefore=6,
    spaceAfter=0,
)
TABLE_TITLE = ParagraphStyle(
    "Table Title",
    fontName="Times-Italic",
    fontSize=12,
    leading=14,
    spaceBefore=0,
    spaceAfter=6,
)
TABLE_NOTE = ParagraphStyle(
    "Table Note",
    fontName="Times-Roman",
    fontSize=10,
    leading=12,
    spaceBefore=4,
    spaceAfter=4,
)


def page_number(canvas, _document) -> None:
    canvas.saveState()
    canvas.setFont("Times-Roman", 12)
    canvas.drawRightString(7.5 * inch, 10.5 * inch, str(canvas.getPageNumber()))
    canvas.restoreState()


def has_page_break(paragraph: DocxParagraph) -> bool:
    return bool(paragraph._p.xpath('.//w:br[@w:type="page"]'))


def paragraph_markup(paragraph: DocxParagraph) -> str:
    parts: list[str] = []
    for run in paragraph.runs:
        text = escape(run.text)
        if not text:
            continue
        if run.bold:
            text = f"<b>{text}</b>"
        if run.italic:
            text = f"<i>{text}</i>"
        parts.append(text)
    return "".join(parts) or escape(paragraph.text)


def iter_blocks(document: Document):
    for child in document.element.body.iterchildren():
        if child.tag.endswith("}p"):
            yield DocxParagraph(child, document)
        elif child.tag.endswith("}tbl"):
            yield DocxTable(child, document)


def title_page_story(document: Document) -> list[object]:
    lines: list[str] = []
    for paragraph in document.paragraphs:
        if has_page_break(paragraph):
            break
        if paragraph.text.strip():
            lines.append(paragraph.text.strip())
    story: list[object] = [Spacer(1, 1.35 * inch)]
    for index, text in enumerate(lines):
        style = TITLE_STYLE if index == 0 else BODY_NO_INDENT
        story.append(Paragraph(escape(text), style))
    story.append(PageBreak())
    return story


def report_table(docx_table: DocxTable) -> Table:
    data: list[list[Paragraph]] = []
    for row_index, row in enumerate(docx_table.rows):
        cells: list[Paragraph] = []
        for cell_index, cell in enumerate(row.cells):
            text = " ".join(p.text.strip() for p in cell.paragraphs if p.text.strip())
            style = ParagraphStyle(
                f"Table cell {row_index}-{cell_index}",
                fontName="Times-Bold" if row_index == 0 else "Times-Roman",
                fontSize=10,
                leading=12,
                alignment=TA_CENTER if cell_index == 1 else TA_LEFT,
            )
            cells.append(Paragraph(escape(text), style))
        data.append(cells)
    table = Table(data, colWidths=[5.0 * inch, 1.5 * inch], repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("LEFTPADDING", (0, 0), (-1, -1), 6),
                ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F2F2F2")),
                ("LINEABOVE", (0, 0), (-1, 0), 0.75, colors.black),
                ("LINEBELOW", (0, 0), (-1, 0), 0.5, colors.black),
                ("LINEBELOW", (0, -1), (-1, -1), 0.75, colors.black),
            ]
        )
    )
    return table


def body_story(document: Document) -> list[object]:
    story: list[object] = []
    passed_title_page = False
    in_references = False
    pending_table_label: Paragraph | None = None
    pending_table_title: Paragraph | None = None

    for block in iter_blocks(document):
        if isinstance(block, DocxParagraph):
            if not passed_title_page:
                if has_page_break(block):
                    passed_title_page = True
                continue
            if has_page_break(block):
                story.append(PageBreak())
                continue
            text = block.text.strip()
            if not text:
                continue
            style_name = block.style.name if block.style is not None else "Normal"
            if text == "References":
                in_references = True
                story.append(Paragraph(escape(text), HEADING_1))
                continue
            if style_name == "Heading 1":
                story.append(Paragraph(escape(text), HEADING_1))
            elif style_name in {"Heading 2", "Heading 3"}:
                story.append(Paragraph(escape(text), HEADING_2))
            elif in_references:
                story.append(Paragraph(paragraph_markup(block), REFERENCE))
            elif block.runs and all(
                (run.font.name or "").lower() == "consolas" for run in block.runs
            ):
                story.append(Paragraph(escape(text), CODE))
            elif text == "Table 1":
                pending_table_label = Paragraph(escape(text), TABLE_LABEL)
            elif text == "Final Raw-Message Test Results":
                pending_table_title = Paragraph(escape(text), TABLE_TITLE)
            elif text.startswith("Note. The final test"):
                story.append(Paragraph(escape(text), TABLE_NOTE))
            elif (
                block.alignment == 1
                and any(bool(run.bold) for run in block.runs)
            ):
                story.append(Paragraph(escape(text), TITLE_STYLE))
            else:
                story.append(Paragraph(escape(text), BODY))
        else:
            table = report_table(block)
            group: list[object] = []
            if pending_table_label is not None:
                group.append(pending_table_label)
                pending_table_label = None
            if pending_table_title is not None:
                group.append(pending_table_title)
                pending_table_title = None
            group.append(table)
            story.append(KeepTogether(group))
    return story


def main() -> int:
    if not DOCX_PATH.exists():
        raise FileNotFoundError("Build the DOCX report before generating the PDF.")
    document = Document(DOCX_PATH)
    story = title_page_story(document) + body_story(document)
    pdf = SimpleDocTemplate(
        str(PDF_PATH),
        pagesize=letter,
        leftMargin=inch,
        rightMargin=inch,
        topMargin=inch,
        bottomMargin=inch,
        title=document.core_properties.title,
        author=document.core_properties.author,
    )
    pdf.build(story, onFirstPage=page_number, onLaterPages=page_number)
    print(f"PDF: {PDF_PATH.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
