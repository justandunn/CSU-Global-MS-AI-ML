"""Build the Module 8 Portfolio Project APA report as a DOCX."""

from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


BASE_DIR = Path(__file__).resolve().parent
OUTPUT_DIR = BASE_DIR / "outputs"
REPORT_DIR = BASE_DIR / "report"
METRICS_PATH = OUTPUT_DIR / "final_metrics.json"
DOCX_PATH = REPORT_DIR / "CSC525_Module8PP_Dunn_Justan.docx"

TITLE = "Option #2: Closed-Domain Customer-Service NLP Chatbot Final Version"
STUDENT = "Justan Dunn"
UNIVERSITY = "Colorado State University Global"
COURSE = "CSC525: Principles of Machine Learning"
INSTRUCTOR = "Dong Nguyen"
DATE = "July 2, 2026"

# Required document design choice with an APA-specific named override.
DESIGN_PRESET = "narrative_proposal"
HEADER_PATTERN = "editorial_cover"
NAMED_OVERRIDE = "CSU_APA7_student_paper"


REFERENCES = [
    (
        "Beaver, I., Freeman, C., & Mueen, A. (2017). An annotated corpus of "
        "relational strategies in customer service. arXiv. "
        "https://doi.org/10.48550/arXiv.1708.05449",
        ["An annotated corpus of relational strategies in customer service"],
    ),
    (
        "Hendrycks, D., & Gimpel, K. (2017). A baseline for detecting "
        "misclassified and out-of-distribution examples in neural networks. "
        "International Conference on Learning Representations. "
        "https://arxiv.org/abs/1610.02136",
        ["International Conference on Learning Representations"],
    ),
    (
        "Manning, C. D., Raghavan, P., & Schutze, H. (2008). Introduction to "
        "information retrieval. Cambridge University Press. "
        "https://nlp.stanford.edu/IR-book/",
        ["Introduction to information retrieval"],
    ),
    (
        "Salton, G., & Buckley, C. (1988). Term-weighting approaches in automatic "
        "text retrieval. Information Processing & Management, 24(5), 513-523. "
        "https://doi.org/10.1016/0306-4573(88)90021-0",
        ["Information Processing & Management", "24"],
    ),
    (
        "Yang, Y., Yih, W.-t., & Meek, C. (2015). WikiQA: A challenge dataset "
        "for open-domain question answering. In L. Marquez, C. Callison-Burch, "
        "& J. Su (Eds.), Proceedings of the 2015 Conference on Empirical Methods "
        "in Natural Language Processing (pp. 2013-2018). Association for "
        "Computational Linguistics. https://doi.org/10.18653/v1/D15-1237",
        ["Proceedings of the 2015 Conference on Empirical Methods in Natural Language Processing"],
    ),
]


def set_run_font(run, *, size: float = 12, bold: bool = False, italic: bool = False) -> None:
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def set_cell_margins(cell, top: int = 80, start: int = 120, bottom: int = 80, end: int = 120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa: list[int]) -> None:
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            set_cell_width(cell, widths_dxa[index])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    set_run_font(run)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instruction, separate, text, end):
        run._r.append(element)


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.5)
    section.footer_distance = Inches(0.5)
    add_page_number(section.header.paragraphs[0])

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal.font.size = Pt(12)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE

    for style_name in ("Heading 1", "Heading 2", "Heading 3"):
        style = document.styles[style_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style.font.size = Pt(12)
        style.font.bold = True
        style.font.color.rgb = None
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(0)
        style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
        style.paragraph_format.keep_with_next = True
    document.styles["Heading 1"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.styles["Heading 2"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    document.styles["Heading 3"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    document.core_properties.title = TITLE
    document.core_properties.author = STUDENT
    document.core_properties.subject = "CSC525 Module 8 Portfolio Project"


def add_centered_line(document: Document, text: str, *, bold: bool = False) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    run = paragraph.add_run(text)
    set_run_font(run, bold=bold)


def add_title_page(document: Document) -> None:
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(72)
    add_centered_line(document, TITLE, bold=True)
    for line in (STUDENT, UNIVERSITY, COURSE, INSTRUCTOR, DATE):
        add_centered_line(document, line)
    document.add_page_break()


def add_body_title(document: Document) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(TITLE)
    set_run_font(run, bold=True)


def add_body_paragraph(document: Document, text: str, *, indent: bool = True) -> None:
    paragraph = document.add_paragraph(style="Normal")
    paragraph.paragraph_format.first_line_indent = Inches(0.5) if indent else None
    paragraph.paragraph_format.keep_together = False
    run = paragraph.add_run(text)
    set_run_font(run)


def add_heading(document: Document, text: str, level: int = 1) -> None:
    paragraph = document.add_paragraph(style=f"Heading {level}")
    paragraph.paragraph_format.first_line_indent = None
    run = paragraph.add_run(text)
    set_run_font(run, bold=True)


def add_code_line(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.5)
    paragraph.paragraph_format.right_indent = Inches(0.5)
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.line_spacing = 1.0
    run = paragraph.add_run(text)
    run.font.name = "Consolas"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Consolas")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Consolas")
    run.font.size = Pt(10)


def add_results_table(document: Document, metrics: dict[str, object]) -> None:
    number = document.add_paragraph()
    number.paragraph_format.space_before = Pt(6)
    number.paragraph_format.space_after = Pt(0)
    number.paragraph_format.line_spacing = 1.0
    set_run_font(number.add_run("Table 1"), bold=True)

    title = document.add_paragraph()
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(6)
    title.paragraph_format.line_spacing = 1.0
    set_run_font(title.add_run("Final Raw-Message Test Results"), italic=True)

    rows = [
        ("Raw route accuracy", metrics["raw_route_accuracy"]),
        ("In-domain coverage", metrics["in_domain_coverage"]),
        ("Accepted route accuracy", metrics["accepted_route_accuracy"]),
        ("Out-of-domain rejection rate", metrics["out_of_domain_rejection_rate"]),
        ("Overall decision accuracy", metrics["overall_decision_accuracy"]),
    ]
    table = document.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "Metric"
    table.rows[0].cells[1].text = "Result"
    header_properties = table.rows[0]._tr.get_or_add_trPr()
    header_marker = OxmlElement("w:tblHeader")
    header_marker.set(qn("w:val"), "1")
    header_properties.append(header_marker)
    for label, value in rows:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = f"{float(value):.4f}"
    set_table_geometry(table, [7200, 2160])

    for row_index, row in enumerate(table.rows):
        for cell_index, cell in enumerate(row.cells):
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.0
                paragraph.alignment = (
                    WD_ALIGN_PARAGRAPH.CENTER
                    if cell_index == 1
                    else WD_ALIGN_PARAGRAPH.LEFT
                )
                for run in paragraph.runs:
                    set_run_font(run, size=10, bold=row_index == 0)

    note = document.add_paragraph()
    note.paragraph_format.space_before = Pt(4)
    note.paragraph_format.space_after = Pt(4)
    note.paragraph_format.line_spacing = 1.0
    set_run_font(
        note.add_run(
            "Note. The final test contains 50 in-domain and 12 out-of-domain raw "
            "messages. No route labels or curated keyword banks were appended to inputs."
        ),
        size=10,
    )


def add_reference(document: Document, text: str, italic_phrases: list[str]) -> None:
    paragraph = document.add_paragraph(style="Normal")
    paragraph.paragraph_format.left_indent = Inches(0.5)
    paragraph.paragraph_format.first_line_indent = Inches(-0.5)
    paragraph.paragraph_format.keep_together = True
    remaining = text
    for phrase in italic_phrases:
        before, separator, after = remaining.partition(phrase)
        if before:
            set_run_font(paragraph.add_run(before))
        if separator:
            set_run_font(paragraph.add_run(separator), italic=True)
            remaining = after
        else:
            remaining = before
    if remaining:
        set_run_font(paragraph.add_run(remaining))


def build_report(metrics_payload: dict[str, object]) -> None:
    test = metrics_payload["test_metrics"]
    calibration = metrics_payload["calibration_metrics"]
    document = Document()
    configure_document(document)
    add_title_page(document)
    add_body_title(document)

    add_body_paragraph(
        document,
        "This portfolio project presents the final version of a closed-domain "
        "customer-service routing and retrieval chatbot. The system accepts a natural-"
        "language support message, estimates the most likely service route, and returns "
        "a controlled response associated with that route. When evidence is weak, it "
        "does not invent an answer; it asks the user to rephrase the request or sends "
        "the case toward human support. The project applies natural language processing "
        "through character n-gram term-frequency inverse-document-frequency (TF-IDF) "
        "features, supervised class centroids, cosine similarity, and calibrated "
        "fallback thresholds. Its purpose is to demonstrate an explainable chatbot that "
        "responds meaningfully while making uncertainty visible."
    )
    add_body_paragraph(
        document,
        "The final design also corrects an important limitation discovered in the "
        "earlier milestones. The Module 5 evaluation produced perfect accuracy when "
        "curated route keywords were appended to held-out questions, but normal users "
        "do not provide those keyword banks. The final model therefore trains on "
        "customer-language questions and is evaluated against separate raw-message "
        "calibration and test sets. This change produces lower but more credible "
        "results and supports the central conclusion of the project: a controlled NLP "
        "chatbot can route common support requests effectively, but it still requires "
        "fallback logic, broader data, and human escalation before enterprise use."
    )

    add_heading(document, "Project Goal and Domain")
    add_body_paragraph(
        document,
        "The chatbot's defined goal is to identify the correct support function for a "
        "customer message and provide an approved response that is clearly related to "
        "the request. It supports ten routes: account access, billing and invoices, "
        "human escalation, onboarding, order status, policy lookup, returns and refunds, "
        "service scheduling, technical support, and training materials. These routes "
        "represent common operational needs that can be adapted to different companies "
        "without exposing actual customer or company data."
    )
    add_body_paragraph(
        document,
        "The chatbot is closed-domain rather than open-domain. This means that it is "
        "designed to answer only within the approved support categories and does not "
        "attempt general conversation, current-events questions, or unrestricted text "
        "generation. A closed-domain design is appropriate because customer-service "
        "answers frequently depend on controlled policies and standard operating "
        "procedures. The system retrieves a route-specific response instead of composing "
        "a new policy. That constraint reduces nonsense responses and creates a clearer "
        "audit trail from an input decision to an approved support action."
    )

    add_heading(document, "Dataset Development and Milestone Continuity")
    add_body_paragraph(
        document,
        "The project developed through four milestones. Module 3 established an "
        "enterprise knowledge-support chatbot using TF-IDF and similarity ranking. "
        "Module 4 refined the use case into customer-service routing and selected a "
        "hybrid data strategy: the RSiCS corpus for relational customer language, "
        "WikiQA for answer relevance and no-answer behavior, and a sanitized FAQ/SOP "
        "corpus for approved domain responses. Beaver et al. (2017) demonstrate that "
        "greetings, backstory, emotion, and other relational language can obscure a "
        "customer's underlying intent. Yang et al. (2015) similarly show the importance "
        "of distinguishing relevant candidate answers from examples for which no "
        "candidate answer is correct. Those findings shaped the final use of neutral "
        "conversational augmentation and explicit out-of-domain tests."
    )
    add_body_paragraph(
        document,
        "Module 5 created 80 sanitized FAQ/SOP questions across the ten routes. Module 6 "
        "then produced a command-line alpha that returned approved answers and used a "
        "similarity threshold for fallback. The final implementation retains that "
        "controlled corpus but separates data roles more rigorously. The 80 original "
        "questions are the only labeled content used for route learning. Each question "
        "is also wrapped in two route-neutral conversational frames, producing 240 "
        "training examples without inserting label names or route-specific keywords. "
        "A separate 52-message calibration set selects the feature configuration and "
        "thresholds. A final 62-message test set contains 50 in-domain messages and 12 "
        "out-of-domain messages. No exact training question appears in either evaluation "
        "set. RSiCS and WikiQA inform the evaluation design but are not represented as "
        "direct training rows in the final local model."
    )

    add_heading(document, "NLP Model, Tools, and Libraries")
    add_body_paragraph(
        document,
        "The final chatbot is implemented in Python 3 with standard-library modules "
        "including csv, json, re, math, argparse, pathlib, and collections. No external "
        "runtime library is required. The trained model is stored as transparent JSON "
        "rather than as an opaque binary object, allowing the vocabulary, inverse "
        "document-frequency values, route centroids, responses, and decision thresholds "
        "to be inspected directly. Supporting scripts reproduce dataset construction, "
        "model selection, evaluation, and validation."
    )
    add_body_paragraph(
        document,
        "TF-IDF represents text by increasing the weight of terms that are informative "
        "within the collection while reducing the influence of terms that appear in many "
        "documents. Salton and Buckley (1988) established term weighting as a strong "
        "baseline for automatic text retrieval, and Manning et al. (2008) explain how "
        "TF-IDF vectors and cosine similarity support ranked retrieval and vector-space "
        "classification. The final implementation applies these ideas to character "
        "n-grams of length three through five. Character n-grams capture word fragments "
        "and are more tolerant of inflection, minor spelling variation, and differences "
        "between words such as invoice, invoicing, and invoices than exact unigram "
        "matching alone."
    )
    add_body_paragraph(
        document,
        "During training, every example is converted to a normalized TF-IDF vector. "
        "Vectors sharing a route label are averaged and normalized to create a supervised "
        "route centroid. At runtime, a raw customer message is transformed with the same "
        "feature configuration and compared with every centroid using cosine similarity. "
        "The highest-scoring route is the proposed intent, and the difference between "
        "the highest and second-highest score is the decision margin. The chatbot answers "
        "only when both the highest score and the margin exceed calibrated thresholds."
    )

    add_heading(document, "Hyperparameter Selection and Fallback Control")
    add_body_paragraph(
        document,
        f"The training process compared four feature configurations: word unigrams, "
        f"word unigrams and bigrams, character 3-5-grams, and a hybrid word-character "
        f"representation. For each configuration, the program evaluated 20 minimum-score "
        f"values and nine minimum-margin values, producing 720 calibration combinations. "
        f"To prevent a model from obtaining a high score by rejecting nearly every "
        f"message, eligible configurations were required to preserve at least 70% "
        f"in-domain coverage and reject at least 60% of out-of-domain calibration inputs. "
        f"The selected configuration was character 3-5-gram TF-IDF, with "
        f"{metrics_payload['vocabulary_size']:,} features, a minimum similarity of "
        f"{float(metrics_payload['minimum_score']):.3f}, and a minimum margin of "
        f"{float(metrics_payload['minimum_margin']):.3f}."
    )
    add_body_paragraph(
        document,
        f"On the calibration set, the selected model achieved an overall decision "
        f"accuracy of {float(calibration['overall_decision_accuracy']):.4f}, an "
        f"in-domain coverage of {float(calibration['in_domain_coverage']):.4f}, and an "
        f"out-of-domain rejection rate of "
        f"{float(calibration['out_of_domain_rejection_rate']):.4f}. Confidence-based "
        f"rejection follows the practical principle that uncertain or out-of-distribution "
        f"inputs should be detectable rather than automatically forced into a known "
        f"class. Hendrycks and Gimpel (2017) describe confidence as a useful baseline "
        f"signal for identifying misclassified and out-of-distribution examples. In this "
        f"project, the score and margin are not treated as probabilities; they are "
        f"decision signals used to choose between a controlled answer and safe fallback."
    )

    add_heading(document, "Final Test Results")
    add_body_paragraph(
        document,
        "Table 1 summarizes the final results. Raw route accuracy measures whether the "
        "highest-scoring centroid matches the expected route before fallback. In-domain "
        "coverage measures how often the chatbot answers instead of rejecting a support "
        "message. Accepted route accuracy measures the correctness of answered in-domain "
        "messages. Out-of-domain rejection measures how often unrelated questions are "
        "sent to fallback. Overall decision accuracy counts a correct accepted route or "
        "a correct out-of-domain rejection as a successful decision."
    )
    add_results_table(document, test)
    add_body_paragraph(
        document,
        f"The final model correctly ranked the route for "
        f"{float(test['raw_route_accuracy']) * 100:.0f}% of in-domain messages. After "
        f"fallback control, it answered {int(test['accepted_in_domain'])} of 50 in-domain "
        f"messages and routed {int(test['correct_accepted_in_domain'])} of those accepted "
        f"messages correctly. It rejected {int(test['rejected_out_of_domain'])} of 12 "
        f"unrelated messages. The remaining errors included "
        f"{int(test['false_reject_in_domain'])} false rejections, "
        f"{int(test['false_accept_out_of_domain'])} false accepts, and three accepted "
        f"misroutes. These errors are retained in the row-level prediction file rather "
        f"than hidden from the submission."
    )
    add_body_paragraph(
        document,
        "The results improve substantially on the fair question-only baseline reported "
        "in Module 5 while avoiding the unsupported claim of perfect performance. The "
        "remaining false accepts show that lexical similarity alone cannot guarantee "
        "domain awareness. For example, a general request that contains words such as "
        "time, solve, or explain can resemble policy or training language even when the "
        "topic is unrelated. A production system would need a larger set of real, "
        "sanitized customer messages, dedicated outlier examples, probability "
        "calibration, monitoring, and periodic human review."
    )

    add_heading(document, "Chatbot Behavior and Running Instructions")
    add_body_paragraph(
        document,
        "The final program provides interactive, single-message, demonstration, and JSON "
        "modes. In interactive mode, the user enters a message, the model predicts a "
        "route, and the program prints the approved response, human-readable route, "
        "confidence label, and similarity score. The commands /topics, /help, and /quit "
        "support navigation. The prepared demonstration contains seven representative "
        "in-domain requests and one unrelated request that activates fallback."
    )
    add_body_paragraph(
        document,
        "The instructor can open the extracted project folder and double-click "
        "run_chatbot.bat on Windows. The same interface can be started from PowerShell "
        "or another terminal with the following command:"
    )
    add_code_line(document, "python chatbot.py --interactive")
    add_body_paragraph(
        document,
        "The prepared demonstration runs with the command below and writes a transcript "
        "to the outputs folder:"
    )
    add_code_line(document, "python chatbot.py --demo")
    add_body_paragraph(
        document,
        "No installation step is required beyond Python 3.10 or newer. The included "
        "README provides additional commands for one-message and JSON output, and the "
        "validation script confirms model, data, report, transcript, and metric files."
    )

    add_heading(document, "Limitations and Responsible Use")
    add_body_paragraph(
        document,
        "The project is a coursework prototype built from synthetic, sanitized examples. "
        "It does not use private customer records, production policies, or current company "
        "systems. Its approved responses are generic workflow guidance rather than actual "
        "service commitments. A company adapting this design would need policy-owner "
        "approval, privacy controls, access management, secure logs, accessibility review, "
        "and testing across the language used by its customers."
    )
    add_body_paragraph(
        document,
        "The classifier also predicts only one primary route. Multi-intent messages, "
        "long conversations, contextual follow-up questions, multilingual inputs, and "
        "changing policies remain outside the final scope. These limits are appropriate "
        "for a small explainable project, but they prevent the model from being treated "
        "as an autonomous customer-service replacement. The fallback response and human "
        "escalation route are therefore core design features rather than exceptions."
    )

    add_heading(document, "Conclusion")
    add_body_paragraph(
        document,
        "The final chatbot satisfies Option #2 by providing a runnable NLP system that "
        "responds to customer-service inputs with relevant, controlled content. It is "
        "closed-domain, dependency-free, reproducible, and documented with complete run "
        "instructions. More importantly, the final evaluation uses raw messages and "
        "reports residual errors instead of relying on metadata-assisted accuracy. The "
        "character n-gram TF-IDF centroid model achieved strong routing performance for "
        "the defined support routes while the calibrated thresholds rejected most "
        "unrelated inputs. The project demonstrates a practical foundation that other "
        "companies could adapt, provided that deployment adds organization-specific data, "
        "governance, security, monitoring, and human oversight."
    )

    add_heading(document, "Use of AI Assistance")
    add_body_paragraph(
        document,
        "AI assistance was used to help structure code, organize project artifacts, and "
        "draft portions of the written report. The scripts were executed locally, the "
        "reported metrics were generated from the included data, and the final package "
        "was validated in the CSC525 workspace. The student remains responsible for "
        "reviewing the submitted work and explaining its methods and limitations."
    )

    document.add_page_break()
    add_heading(document, "References")
    for reference, italic_phrases in REFERENCES:
        add_reference(document, reference, italic_phrases)

    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    document.save(DOCX_PATH)
    print(f"Design preset: {DESIGN_PRESET}")
    print(f"Header pattern: {HEADER_PATTERN}")
    print(f"Named override: {NAMED_OVERRIDE}")
    print(f"DOCX: {DOCX_PATH.resolve()}")


def main() -> int:
    if not METRICS_PATH.exists():
        raise FileNotFoundError("Run train_final_chatbot.py before building the report.")
    metrics = json.loads(METRICS_PATH.read_text(encoding="utf-8"))
    build_report(metrics)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
