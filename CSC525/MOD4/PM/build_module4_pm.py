"""Build the Module 4 Portfolio Milestone paper."""

from __future__ import annotations

from html import escape
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer


OUTPUT_DIR = Path(__file__).resolve().parent
DOCX_PATH = OUTPUT_DIR / "CSC525_Module4PM_Dunn_Justan.docx"
PDF_PATH = OUTPUT_DIR / "CSC525_Module4PM_Dunn_Justan.pdf"

TITLE = "Option #2: NLP Chatbot Project Data Collection"
RUNNING_TITLE = "NLP Chatbot Data Collection"
STUDENT_NAME = "Justan Dunn"
UNIVERSITY = "Colorado State University Global"
COURSE = "CSC525: Principles of Machine Learning"
ASSIGNMENT = "Module 4 Portfolio Milestone"
INSTRUCTOR = "Dong Nguyen"
DATE = "June 7, 2026"


SECTIONS = [
    (
        None,
        [
            (
                "For the portfolio project, the selected path is Option #2, the NLP Chatbot Project. "
                "The proposed system is a closed-domain customer-service routing and support chatbot "
                "that could be adapted for companies in different industries. The chatbot's business "
                "purpose is to interpret a customer's natural-language message, identify the most likely "
                "service category, retrieve an approved response when one is available, and route low-"
                "confidence cases to a human agent. This milestone focuses on the data collection and "
                "training strategy because the value of the chatbot will depend less on a complicated "
                "interface and more on whether its data represents real customer language, reliable "
                "answer selection, and the specific policies or procedures of the organization using it."
            ),
            (
                "The most appropriate data strategy is a hybrid approach that combines public research "
                "datasets with a small, controlled enterprise knowledge base. A public dataset can help "
                "the model learn general customer-service language patterns, but a company-specific "
                "dataset is still needed because customer-support answers must be tied to approved "
                "business rules. For that reason, the final project should not rely on an open-domain "
                "chat dataset such as movie dialogue or broad social media posts. Those datasets may be "
                "interesting for general conversation, but they are a weak fit for a business chatbot "
                "that needs accurate routing, traceable answers, and safe fallback behavior."
            ),
        ],
    ),
    (
        "Dataset Selection Strategy",
        [
            (
                "The first selected dataset is the Relational Strategies in Customer Service corpus. "
                "Beaver et al. (2017) describe this dataset as a publicly available commercial customer-"
                "service corpus built from human interactions with intelligent virtual agents in travel "
                "and telecommunications. The dataset is useful for this project because customer messages "
                "often contain greetings, backstory, emotion, gratitude, frustration, and other relational "
                "language that may not directly identify the user's intent. For a routing chatbot, those "
                "segments matter because the model must recognize the core request even when it is mixed "
                "with conversational text. RSiCS can therefore support preprocessing rules, intent-labeling "
                "examples, and tests for whether the chatbot can separate the customer's service need from "
                "the extra language around it."
            ),
            (
                "The second selected dataset is the Microsoft Research WikiQA Corpus. WikiQA contains "
                "question and candidate-answer sentence pairs, including answer labels and examples where "
                "no candidate sentence is correct (Yang et al., 2015). This is important for the portfolio "
                "project because the chatbot should not always assume that the best available match is good "
                "enough to send to a customer. WikiQA can be used to train and test answer-selection logic "
                "by teaching the model to rank candidate responses and by supporting an answer-triggering "
                "threshold. In practical terms, WikiQA helps with the retrieval side of the chatbot: given "
                "a customer question, the model should identify whether a candidate answer is relevant "
                "enough to return."
            ),
            (
                "The third dataset will be a custom, sanitized enterprise FAQ and standard operating "
                "procedure corpus. This corpus should contain approved categories such as billing, account "
                "access, order status, product troubleshooting, returns, service scheduling, and escalation. "
                "Each row should include a route label, sample customer question, approved answer, source "
                "document or policy owner, and optional keywords. This controlled dataset is necessary "
                "because public data can teach general language behavior, but it cannot safely define a "
                "company's actual return policy, account rules, service-level commitments, or escalation "
                "processes. The custom corpus also gives the final project a clear audit trail from each "
                "answer back to the business source that approved it."
            ),
        ],
    ),
    (
        "Training Plan",
        [
            (
                "The initial model should be retrieval-based rather than generative. A retrieval-based "
                "design is more appropriate for the course project because it is easier to explain, test, "
                "and control. The first version will use Python, pandas, and scikit-learn. Text will be "
                "normalized by lowercasing, trimming whitespace, removing unnecessary punctuation, and "
                "combining the sample question, keywords, and route label into a searchable training field. "
                "The model will then use a term frequency-inverse document frequency representation to "
                "convert the training text into weighted features. This design follows a standard "
                "information-retrieval approach in which documents and queries are represented in a vector "
                "space and then ranked by similarity (Manning et al., 2008)."
            ),
            (
                "In implementation, scikit-learn's TfidfVectorizer can convert the approved question and "
                "answer entries into a TF-IDF document-term matrix, and the fitted vectorizer can transform "
                "new customer inputs into the same feature space (scikit-learn developers, n.d.). The "
                "main tuning choices will include unigram and bigram features, document-frequency limits "
                "to reduce overly common terms, and a minimum similarity threshold for returning an answer. "
                "WikiQA will be used to test whether the retrieval logic can rank relevant answer sentences "
                "above irrelevant candidates. RSiCS will be used to test whether preprocessing and routing "
                "still work when customer messages include conversational or emotional language. The custom "
                "FAQ/SOP corpus will be split into training and validation examples so the final chatbot can "
                "be evaluated on company-specific questions it did not directly memorize."
            ),
            (
                "Evaluation should include top-1 accuracy, top-3 accuracy, route-label accuracy, and "
                "fallback quality. Top-1 accuracy checks whether the first answer is correct. Top-3 "
                "accuracy checks whether the correct response appears among the best few candidates, which "
                "is helpful if the interface later asks the customer to choose from suggested topics. "
                "Fallback quality checks whether the chatbot refuses to answer when similarity is too low. "
                "This is a key technical consideration because an enterprise customer-service chatbot "
                "should make uncertainty visible instead of producing unsupported answers."
            ),
        ],
    ),
    (
        "Chatbot Use of the Training",
        [
            (
                "At runtime, the chatbot will apply the same preprocessing steps used during training. "
                "A customer message will be normalized, transformed into the fitted TF-IDF vector space, "
                "and compared with the stored FAQ/SOP entries. If the best match exceeds the confidence "
                "threshold, the bot will return the approved answer and route label. If the score is "
                "moderate, the bot can ask a clarifying question or present the closest few categories. "
                "If the score is low, the bot will route the conversation to a human agent or provide a "
                "safe fallback such as asking the customer to rephrase the request. This workflow lets "
                "the chatbot use training data to control both answer selection and escalation behavior."
            ),
            (
                "The design also supports later improvement. Logged customer questions can be reviewed, "
                "deduplicated, sanitized, and added to the custom corpus only after approval. That process "
                "prevents the model from learning private or incorrect responses directly from production "
                "traffic. It also gives the company a practical feedback loop: common fallback questions "
                "become candidates for new FAQ entries, confusing categories can be rewritten, and weak "
                "routing labels can be corrected. This makes the chatbot useful as a managed business "
                "system rather than a one-time class demonstration."
            ),
        ],
    ),
    (
        "Conclusion",
        [
            (
                "The recommended dataset plan for the Module 8 portfolio project is to use RSiCS for "
                "customer-service language patterns, WikiQA for answer-selection and threshold training, "
                "and a sanitized enterprise FAQ/SOP corpus for the final domain-specific knowledge base. "
                "This combination meets the technical needs of a practical customer-service chatbot while "
                "keeping the project explainable, testable, and appropriate for business use. Most "
                "importantly, it separates general language learning from approved company answers, which "
                "is necessary for a chatbot that could be implemented safely across different organizations."
            ),
        ],
    ),
]


REFERENCES = [
    (
        "Beaver, I., Freeman, C., & Mueen, A. (2017). An annotated corpus of relational strategies "
        "in customer service. arXiv. https://doi.org/10.48550/arXiv.1708.05449"
    ),
    (
        "Manning, C. D., Raghavan, P., & Schutze, H. (2008). Introduction to information retrieval. "
        "Cambridge University Press. https://nlp.stanford.edu/IR-book/"
    ),
    (
        "scikit-learn developers. (n.d.). TfidfVectorizer. scikit-learn documentation. "
        "https://scikit-learn.org/stable/modules/generated/sklearn.feature_extraction.text.TfidfVectorizer.html"
    ),
    (
        "Yang, Y., Yih, W.-t., & Meek, C. (2015). WikiQA: A challenge dataset for open-domain "
        "question answering. In L. Marquez, C. Callison-Burch, & J. Su (Eds.), Proceedings of the "
        "2015 Conference on Empirical Methods in Natural Language Processing (pp. 2013-2018). "
        "Association for Computational Linguistics. https://doi.org/10.18653/v1/D15-1237"
    ),
]


def add_page_number(paragraph) -> None:
    """Add a Word PAGE field to a header paragraph."""
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()

    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")

    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"

    fld_char_separate = OxmlElement("w:fldChar")
    fld_char_separate.set(qn("w:fldCharType"), "separate")

    text = OxmlElement("w:t")
    text.text = "1"

    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")

    run._r.append(fld_char_begin)
    run._r.append(instr_text)
    run._r.append(fld_char_separate)
    run._r.append(text)
    run._r.append(fld_char_end)


def set_run_font(run, *, bold: bool = False, italic: bool = False) -> None:
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(12)
    run.bold = bold
    run.italic = italic


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.5)
    section.footer_distance = Inches(0.5)

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 2.0
    normal.paragraph_format.space_after = Pt(0)

    for style_name in ("Title", "Heading 1", "Heading 2"):
        style = document.styles[style_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.size = Pt(12)
        style.font.bold = True
        style.font.color.rgb = None
        style.paragraph_format.line_spacing = 2.0
        style.paragraph_format.space_after = Pt(0)

    add_page_number(section.header.paragraphs[0])


def add_centered_line(document: Document, text: str, *, bold: bool = False) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.line_spacing = 2.0
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    set_run_font(run, bold=bold)


def add_body_paragraph(document: Document, text: str, *, indent: bool = True) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.line_spacing = 2.0
    paragraph.paragraph_format.space_after = Pt(0)
    if indent:
        paragraph.paragraph_format.first_line_indent = Inches(0.5)
    run = paragraph.add_run(text)
    set_run_font(run)


def add_section_heading(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.line_spacing = 2.0
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    set_run_font(run, bold=True)


def add_reference(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.line_spacing = 2.0
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.left_indent = Inches(0.5)
    paragraph.paragraph_format.first_line_indent = Inches(-0.5)
    run = paragraph.add_run(text)
    set_run_font(run)


def build_docx() -> None:
    document = Document()
    configure_document(document)

    for _ in range(4):
        document.add_paragraph()
    add_centered_line(document, TITLE, bold=True)
    document.add_paragraph()
    add_centered_line(document, STUDENT_NAME)
    add_centered_line(document, UNIVERSITY)
    add_centered_line(document, COURSE)
    add_centered_line(document, ASSIGNMENT)
    add_centered_line(document, INSTRUCTOR)
    add_centered_line(document, DATE)

    document.add_page_break()

    add_section_heading(document, TITLE)
    for heading, paragraphs in SECTIONS:
        if heading is not None:
            add_section_heading(document, heading)
        for paragraph in paragraphs:
            add_body_paragraph(document, paragraph)

    document.add_page_break()
    add_section_heading(document, "References")
    for reference in REFERENCES:
        add_reference(document, reference)

    document.save(DOCX_PATH)
    print(DOCX_PATH)


def on_page(canvas, _doc) -> None:
    canvas.saveState()
    canvas.setFont("Times-Roman", 12)
    canvas.drawRightString(7.5 * inch, 10.5 * inch, str(canvas.getPageNumber()))
    canvas.restoreState()


def pdf_paragraph(text: str, style: ParagraphStyle) -> Paragraph:
    return Paragraph(escape(text), style)


def build_pdf() -> None:
    document = SimpleDocTemplate(
        str(PDF_PATH),
        pagesize=letter,
        rightMargin=inch,
        leftMargin=inch,
        topMargin=inch,
        bottomMargin=inch,
    )

    title_style = ParagraphStyle(
        "TitleStyle",
        fontName="Times-Bold",
        fontSize=12,
        leading=24,
        alignment=TA_CENTER,
        spaceAfter=0,
    )
    center_style = ParagraphStyle(
        "CenterStyle",
        fontName="Times-Roman",
        fontSize=12,
        leading=24,
        alignment=TA_CENTER,
        spaceAfter=0,
    )
    body_style = ParagraphStyle(
        "BodyStyle",
        fontName="Times-Roman",
        fontSize=12,
        leading=24,
        firstLineIndent=0.5 * inch,
        alignment=TA_LEFT,
        spaceAfter=0,
    )
    heading_style = ParagraphStyle(
        "HeadingStyle",
        fontName="Times-Bold",
        fontSize=12,
        leading=24,
        alignment=TA_CENTER,
        spaceAfter=0,
    )
    refs_style = ParagraphStyle(
        "ReferencesStyle",
        fontName="Times-Roman",
        fontSize=12,
        leading=24,
        leftIndent=0.5 * inch,
        firstLineIndent=-0.5 * inch,
        alignment=TA_LEFT,
        spaceAfter=0,
    )

    story = []
    story.append(Spacer(1, 1.35 * inch))
    story.append(pdf_paragraph(TITLE, title_style))
    story.append(Spacer(1, 0.24 * inch))
    for line in (STUDENT_NAME, UNIVERSITY, COURSE, ASSIGNMENT, INSTRUCTOR, DATE):
        story.append(pdf_paragraph(line, center_style))

    story.append(PageBreak())
    story.append(pdf_paragraph(TITLE, title_style))
    for heading, paragraphs in SECTIONS:
        if heading is not None:
            story.append(pdf_paragraph(heading, heading_style))
        for paragraph in paragraphs:
            story.append(pdf_paragraph(paragraph, body_style))

    story.append(PageBreak())
    story.append(pdf_paragraph("References", title_style))
    for reference in REFERENCES:
        story.append(pdf_paragraph(reference, refs_style))

    document.build(story, onFirstPage=on_page, onLaterPages=on_page)
    print(PDF_PATH)


def main() -> None:
    build_docx()
    build_pdf()


if __name__ == "__main__":
    main()
