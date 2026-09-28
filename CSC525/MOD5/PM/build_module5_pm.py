"""Build the Module 5 Portfolio Milestone report as DOCX and PDF."""

from __future__ import annotations

import json
import csv
from html import escape
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer


OUTPUT_DIR = Path(__file__).resolve().parent
REPORT_DIR = OUTPUT_DIR / "report"
METRICS_PATH = OUTPUT_DIR / "outputs" / "training_metrics.json"
SWEEP_PATH = OUTPUT_DIR / "outputs" / "hyperparameter_sweep_summary.csv"
DOCX_PATH = REPORT_DIR / "CSC525_Module5PM_Dunn_Justan.docx"
PDF_PATH = REPORT_DIR / "CSC525_Module5PM_Dunn_Justan.pdf"

TITLE = "Option #2: NLP Chatbot Project Training"
STUDENT_NAME = "Justan Dunn"
UNIVERSITY = "Colorado State University Global"
COURSE = "CSC525: Principles of Machine Learning"
ASSIGNMENT = "Module 5 Portfolio Milestone"
INSTRUCTOR = "Dong Nguyen"
DATE = "June 21, 2026"


REFERENCES = [
    (
        "Beaver, I., Freeman, C., & Mueen, A. (2017). An annotated corpus of "
        "relational strategies in customer service. arXiv. "
        "https://doi.org/10.48550/arXiv.1708.05449"
    ),
    (
        "Manning, C. D., Raghavan, P., & Schutze, H. (2008). Introduction to "
        "information retrieval. Cambridge University Press. "
        "https://nlp.stanford.edu/IR-book/"
    ),
    (
        "Yang, Y., Yih, W.-t., & Meek, C. (2015). WikiQA: A challenge dataset "
        "for open-domain question answering. In L. Marquez, C. Callison-Burch, "
        "& J. Su (Eds.), Proceedings of the 2015 Conference on Empirical "
        "Methods in Natural Language Processing (pp. 2013-2018). Association "
        "for Computational Linguistics. https://doi.org/10.18653/v1/D15-1237"
    ),
]


def load_metrics() -> dict:
    metrics = json.loads(METRICS_PATH.read_text(encoding="utf-8"))
    if SWEEP_PATH.exists():
        with SWEEP_PATH.open("r", encoding="utf-8", newline="") as file_obj:
            metrics["sweep_rows"] = list(csv.DictReader(file_obj))
    else:
        metrics["sweep_rows"] = []
    return metrics


def body_sections(metrics: dict) -> list[tuple[str | None, list[str]]]:
    values = metrics["metrics"]
    hyperparameters = metrics["hyperparameters"]
    train_rows = metrics["train_rows"]
    test_rows = metrics["test_rows"]
    distinct_routes = metrics["distinct_routes"]
    vocab_size = metrics["vocabulary_size"]
    accuracy = values["accuracy"]
    macro_precision = values["macro_precision"]
    macro_recall = values["macro_recall"]
    macro_f1 = values["macro_f1"]
    mean_similarity = values["mean_similarity"]
    seed = hyperparameters["random_seed"]
    ratio = hyperparameters["test_group_ratio"]
    ngram_range = hyperparameters["ngram_range"]
    max_features = hyperparameters["max_features"]
    threshold = hyperparameters["similarity_threshold"]
    text_column = metrics["training_text_column"]

    question_only_rows = [
        row for row in metrics.get("sweep_rows", [])
        if row.get("text_column") == "sample_question"
    ]
    best_question_only = max(
        question_only_rows,
        key=lambda row: float(row.get("accuracy", 0)),
        default={"accuracy": "0.0000", "macro_f1": "0.0000"},
    )

    return [
        (
            None,
            [
                (
                    "For the Module 5 Portfolio Milestone, the selected project path remains "
                    "Option #2, the NLP Chatbot Project. This revised milestone follows the "
                    "dataset strategy selected in the Module 4 Portfolio Milestone: RSiCS for "
                    "customer-service language patterns, WikiQA for answer-selection and "
                    "threshold testing, and a sanitized enterprise FAQ/SOP corpus for the "
                    "domain-specific knowledge base. The Module 5 training run focused on the "
                    "custom FAQ/SOP corpus because it is the part of the data strategy that "
                    "contains approved route labels, knowledge-base wording, and response text."
                ),
            ],
        ),
        (
            "Model Type and Dataset",
            [
                (
                    "The first trained model was a retrieval-based TF-IDF cosine-similarity "
                    "baseline. This model type is a better match for the Module 4 plan than the "
                    "previous classifier because the proposed chatbot is intended to retrieve or "
                    "route to approved knowledge-base content instead of generating open-ended "
                    "answers. TF-IDF converts the FAQ/SOP entries into weighted term vectors, "
                    "and cosine similarity ranks the closest approved entry for a new support "
                    "message. This follows the general information-retrieval approach described "
                    "by Manning et al. (2008)."
                ),
                (
                    "The actual training file was a sanitized enterprise FAQ/SOP corpus created "
                    "for this coursework prototype. It contains 80 entries across 10 route "
                    "labels, including account access, billing and invoice, order status, "
                    "returns and refunds, service scheduling, technical support, policy lookup, "
                    "onboarding, training materials, and human escalation. Each row includes a "
                    "sample question, keywords, approved answer, category, route label, and a "
                    "source note. RSiCS and WikiQA were not used as direct training inputs in "
                    "this first run; instead, they remain the planned supplemental datasets. "
                    "RSiCS can later test conversational customer-service language variation "
                    "(Beaver et al., 2017), while WikiQA can later test answer ranking and "
                    "no-answer threshold behavior (Yang et al., 2015)."
                ),
            ],
        ),
        (
            "Hyperparameters and Training Setup",
            [
                (
                    f"The final training run used {train_rows} training rows and {test_rows} "
                    f"test rows across {distinct_routes} route labels. The split was grouped by "
                    "source row so that each FAQ/SOP entry remained in either training or "
                    "testing. The held-out group ratio was "
                    f"{ratio:.2f}, and the random seed was {seed}. This keeps the evaluation "
                    "reproducible while still holding back examples from every route."
                ),
                (
                    f"The final representation used the {text_column} field, lowercase "
                    f"{ngram_range[0]}-{ngram_range[1]} n-gram TF-IDF vectors, and a vocabulary "
                    f"of {vocab_size} features. The maximum feature limit was {max_features}, "
                    f"and the draft similarity threshold was {threshold:.2f}. The selected text "
                    "field combines the user's sample question with curated keywords because the "
                    "Module 4 plan framed the custom dataset as a managed FAQ/SOP knowledge base, "
                    "not as unlabeled raw chat logs."
                ),
            ],
        ),
        (
            "Training Results",
            [
                (
                    f"The first formal training run produced an accuracy of {accuracy:.4f}, "
                    f"macro precision of {macro_precision:.4f}, macro recall of "
                    f"{macro_recall:.4f}, macro F1 score of {macro_f1:.4f}, and mean similarity "
                    f"of {mean_similarity:.4f}. No test entry fell below the 0.20 similarity "
                    "threshold. This strong result shows that the retrieval model can separate "
                    "the controlled FAQ/SOP routes when each entry includes approved keywords "
                    "and domain-specific support language."
                ),
                (
                    "The result should be interpreted carefully. A hyperparameter sweep showed "
                    "that the best question-only run reached only "
                    f"{float(best_question_only['accuracy']):.4f} accuracy and "
                    f"{float(best_question_only['macro_f1']):.4f} macro F1. That contrast is "
                    "important: the controlled FAQ/SOP metadata made the first retrieval model "
                    "work, but raw customer language still needs more examples and broader "
                    "variation. The high score is therefore a proof that the knowledge-base "
                    "structure is usable, not proof that the chatbot is production-ready."
                ),
            ],
        ),
        (
            "Reflection and Next Steps",
            [
                (
                    "This redo is more consistent with the prior portfolio plan because it trains "
                    "against the intended enterprise knowledge-base format. The training experience "
                    "also clarified a key project risk: a retrieval chatbot depends heavily on the "
                    "quality of its curated entries. When route labels, approved answers, and "
                    "keywords are clean, the model can route reliably. When only short customer "
                    "questions are available, the same small corpus does not provide enough lexical "
                    "coverage for dependable routing."
                ),
                (
                    "The next step is to expand the custom corpus and then test it against the two "
                    "public datasets selected in Module 4. RSiCS should be used to add realistic "
                    "customer-service phrasing, including relational or emotional language that may "
                    "surround the core request. WikiQA should be used to test whether a candidate "
                    "answer is relevant enough to return or whether the chatbot should ask a "
                    "clarifying question. This preserves the business-safe design: high-confidence "
                    "matches return approved content, uncertain matches ask for clarification, and "
                    "low-confidence cases escalate to a human agent."
                ),
            ],
        ),
        (
            "Use of AI Assistance",
            [
                (
                    "AI assistance was used to help structure the training script, organize the "
                    "milestone files, and draft this report. The model was executed locally, and "
                    "the reported metrics, validation checks, and output files were generated from "
                    "the local CSC525 workspace before submission."
                ),
            ],
        ),
    ]


def set_run_font(run, *, bold: bool = False, italic: bool = False) -> None:
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(12)
    run.bold = bold
    run.italic = italic


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()

    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char_separate = OxmlElement("w:fldChar")
    fld_char_separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")

    run._r.append(fld_char_begin)
    run._r.append(instr_text)
    run._r.append(fld_char_separate)
    run._r.append(text)
    run._r.append(fld_char_end)


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.5)
    section.footer_distance = Inches(0.5)

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 2.0
    normal.paragraph_format.space_after = Pt(0)

    for style_name in ("Heading 1", "Heading 2"):
        style = document.styles[style_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.size = Pt(12)
        style.font.bold = True
        style.paragraph_format.line_spacing = 2.0
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(0)

    add_page_number(section.header.paragraphs[0])


def add_centered_line(document: Document, text: str, *, bold: bool = False) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.line_spacing = 2.0
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    set_run_font(run, bold=bold)


def add_body_paragraph(document: Document, text: str, *, indent: bool = True) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.line_spacing = 2.0
    paragraph.paragraph_format.space_after = Pt(0)
    if indent:
        paragraph.paragraph_format.first_line_indent = Inches(0.5)
    run = paragraph.add_run(text)
    set_run_font(run)


def add_section_heading(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.line_spacing = 2.0
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(text)
    set_run_font(run, bold=True)


def add_reference(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.line_spacing = 2.0
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.left_indent = Inches(0.5)
    paragraph.paragraph_format.first_line_indent = Inches(-0.5)
    run = paragraph.add_run(text)
    set_run_font(run)


def build_docx(metrics: dict) -> None:
    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    document = Document()
    configure_document(document)

    for _ in range(4):
        document.add_paragraph()
    add_centered_line(document, TITLE, bold=True)
    document.add_paragraph()
    for line in (STUDENT_NAME, UNIVERSITY, COURSE, ASSIGNMENT, INSTRUCTOR, DATE):
        add_centered_line(document, line)

    document.add_page_break()
    add_centered_line(document, TITLE, bold=True)
    for heading, paragraphs in body_sections(metrics):
        if heading is not None:
            if heading in {"Hyperparameters and Training Setup", "Reflection and Next Steps"}:
                document.add_page_break()
            add_section_heading(document, heading)
        for paragraph in paragraphs:
            add_body_paragraph(document, paragraph)

    document.add_page_break()
    add_centered_line(document, "References", bold=True)
    for reference in REFERENCES:
        add_reference(document, reference)

    document.save(DOCX_PATH)
    print(DOCX_PATH)


def pdf_paragraph(text: str, style: ParagraphStyle) -> Paragraph:
    return Paragraph(escape(text), style)


def on_page(canvas, _doc) -> None:
    canvas.saveState()
    canvas.setFont("Times-Roman", 12)
    canvas.drawRightString(7.5 * inch, 10.5 * inch, str(canvas.getPageNumber()))
    canvas.restoreState()


def build_pdf(metrics: dict) -> None:
    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    document = SimpleDocTemplate(
        str(PDF_PATH),
        pagesize=letter,
        rightMargin=inch,
        leftMargin=inch,
        topMargin=inch,
        bottomMargin=inch,
    )

    title_style = ParagraphStyle(
        "TitleStyle",
        fontName="Times-Bold",
        fontSize=12,
        leading=24,
        alignment=TA_CENTER,
        spaceAfter=0,
    )
    center_style = ParagraphStyle(
        "CenterStyle",
        fontName="Times-Roman",
        fontSize=12,
        leading=24,
        alignment=TA_CENTER,
        spaceAfter=0,
    )
    body_style = ParagraphStyle(
        "BodyStyle",
        fontName="Times-Roman",
        fontSize=12,
        leading=24,
        firstLineIndent=0.5 * inch,
        alignment=TA_LEFT,
        spaceAfter=0,
    )
    heading_style = ParagraphStyle(
        "HeadingStyle",
        fontName="Times-Bold",
        fontSize=12,
        leading=24,
        alignment=TA_LEFT,
        spaceAfter=0,
    )
    refs_style = ParagraphStyle(
        "ReferencesStyle",
        fontName="Times-Roman",
        fontSize=12,
        leading=24,
        leftIndent=0.5 * inch,
        firstLineIndent=-0.5 * inch,
        alignment=TA_LEFT,
        spaceAfter=0,
    )

    story = [Spacer(1, 1.35 * inch), pdf_paragraph(TITLE, title_style), Spacer(1, 0.24 * inch)]
    for line in (STUDENT_NAME, UNIVERSITY, COURSE, ASSIGNMENT, INSTRUCTOR, DATE):
        story.append(pdf_paragraph(line, center_style))

    story.append(PageBreak())
    story.append(pdf_paragraph(TITLE, title_style))
    for heading, paragraphs in body_sections(metrics):
        if heading is not None:
            if heading in {"Hyperparameters and Training Setup", "Reflection and Next Steps"}:
                story.append(PageBreak())
            story.append(pdf_paragraph(heading, heading_style))
        for paragraph in paragraphs:
            story.append(pdf_paragraph(paragraph, body_style))

    story.append(PageBreak())
    story.append(pdf_paragraph("References", title_style))
    for reference in REFERENCES:
        story.append(pdf_paragraph(reference, refs_style))

    document.build(story, onFirstPage=on_page, onLaterPages=on_page)
    print(PDF_PATH)


def main() -> None:
    metrics = load_metrics()
    build_docx(metrics)
    build_pdf(metrics)


if __name__ == "__main__":
    main()
