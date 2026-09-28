"""Build the Module 6 Portfolio Milestone report as DOCX and PDF."""

from __future__ import annotations

import csv
import json
from html import escape
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer


BASE_DIR = Path(__file__).resolve().parent
COURSE_DIR = BASE_DIR.parents[1]
REPORT_DIR = BASE_DIR / "report"
OUTPUT_DIR = BASE_DIR / "outputs"
MODULE5_METRICS_PATH = (
    COURSE_DIR / "MOD5" / "PM" / "outputs" / "training_metrics.json"
)
LOCAL_METRICS_SNAPSHOT = OUTPUT_DIR / "training_metrics_snapshot.json"
DEMO_PREDICTIONS_PATH = OUTPUT_DIR / "alpha_demo_predictions.csv"
DOCX_PATH = REPORT_DIR / "CSC525_Module6PM_Dunn_Justan.docx"
PDF_PATH = REPORT_DIR / "CSC525_Module6PM_Dunn_Justan.pdf"

TITLE = "Option #2: NLP Chatbot Project Alpha"
STUDENT_NAME = "Justan Dunn"
UNIVERSITY = "Colorado State University Global"
COURSE = "CSC525: Principles of Machine Learning"
ASSIGNMENT = "Module 6 Portfolio Milestone"
INSTRUCTOR = "Dong Nguyen"
DATE = "June 21, 2026"

REFERENCES = [
    (
        "Beaver, I., Freeman, C., & Mueen, A. (2017). An annotated corpus of "
        "relational strategies in customer service. arXiv. "
        "https://doi.org/10.48550/arXiv.1708.05449"
    ),
    (
        "Manning, C. D., Raghavan, P., & Schutze, H. (2008). Introduction to "
        "information retrieval. Cambridge University Press. "
        "https://nlp.stanford.edu/IR-book/"
    ),
    (
        "Pedregosa, F., Varoquaux, G., Gramfort, A., Michel, V., Thirion, B., "
        "Grisel, O., Blondel, M., Prettenhofer, P., Weiss, R., Dubourg, V., "
        "Vanderplas, J., Passos, A., Cournapeau, D., Brucher, M., Perrot, M., "
        "& Duchesnay, E. (2011). Scikit-learn: Machine learning in Python. "
        "Journal of Machine Learning Research, 12, 2825-2830. "
        "https://jmlr.org/papers/v12/pedregosa11a.html"
    ),
    (
        "Yang, Y., Yih, W.-t., & Meek, C. (2015). WikiQA: A challenge dataset "
        "for open-domain question answering. In L. Marquez, C. Callison-Burch, "
        "& J. Su (Eds.), Proceedings of the 2015 Conference on Empirical "
        "Methods in Natural Language Processing (pp. 2013-2018). Association "
        "for Computational Linguistics. https://doi.org/10.18653/v1/D15-1237"
    ),
]


def load_metrics() -> dict[str, Any]:
    if MODULE5_METRICS_PATH.exists():
        metrics = json.loads(MODULE5_METRICS_PATH.read_text(encoding="utf-8"))
    elif LOCAL_METRICS_SNAPSHOT.exists():
        metrics = json.loads(LOCAL_METRICS_SNAPSHOT.read_text(encoding="utf-8"))
    else:
        raise FileNotFoundError(
            "Could not find Module 5 training metrics or local metrics snapshot."
        )
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    LOCAL_METRICS_SNAPSHOT.write_text(json.dumps(metrics, indent=2), encoding="utf-8")
    return metrics


def load_demo_predictions() -> list[dict[str, str]]:
    if not DEMO_PREDICTIONS_PATH.exists():
        raise FileNotFoundError(
            "Run chatbot_alpha.py --demo before building the Module 6 report."
        )
    with DEMO_PREDICTIONS_PATH.open("r", encoding="utf-8", newline="") as file_obj:
        return list(csv.DictReader(file_obj))


def project_facts(metrics: dict[str, Any], predictions: list[dict[str, str]]) -> dict[str, Any]:
    accepted = [
        row
        for row in predictions
        if str(row.get("above_threshold", "")).lower() == "true"
    ]
    fallback = len(predictions) - len(accepted)
    return {
        "train_rows": metrics["train_rows"],
        "test_rows": metrics["test_rows"],
        "distinct_routes": metrics["distinct_routes"],
        "vocabulary_size": metrics["vocabulary_size"],
        "accuracy": metrics["metrics"]["accuracy"],
        "macro_f1": metrics["metrics"]["macro_f1"],
        "mean_similarity": metrics["metrics"]["mean_similarity"],
        "module5_threshold": metrics["hyperparameters"]["similarity_threshold"],
        "demo_exchanges": len(predictions),
        "accepted": len(accepted),
        "fallback": fallback,
        "alpha_threshold": 0.35,
    }


def body_sections(facts: dict[str, Any]) -> list[tuple[str | None, list[str]]]:
    return [
        (
            None,
            [
                (
                    "The selected portfolio path remains Option #2, the NLP Chatbot "
                    "Project. The current alpha is a closed-domain customer-service "
                    "routing and retrieval chatbot that uses a sanitized enterprise "
                    "FAQ/SOP corpus. This direction is consistent with the Module 4 "
                    "dataset plan and the Module 5 training milestone: public datasets "
                    "will be used for broader validation, but the final chatbot should "
                    "return controlled support content rather than open-ended generated "
                    "responses."
                ),
                (
                    "The purpose of the alpha is to prove that the basic retrieval "
                    "workflow can accept a natural-language support message, identify a "
                    "likely support route, and return an approved response. The final "
                    "Module 8 version must go further by giving the instructor a runnable "
                    "or accessible chatbot, explaining the model and tools, and showing "
                    "that the responses are clearly related to user inputs."
                ),
            ],
        ),
        (
            "Current Alpha Functionality",
            [
                (
                    "The alpha chatbot is implemented as a local Python command-line "
                    "program. It loads the custom FAQ/SOP corpus, builds lowercase "
                    "unigram TF-IDF vectors, and uses cosine similarity to compare a "
                    "new user message against approved knowledge-base entries. This is "
                    "a standard information-retrieval design because term weighting and "
                    "cosine ranking are well suited to matching short queries to "
                    "candidate documents (Manning et al., 2008)."
                ),
                (
                    f"The current corpus contains {facts['distinct_routes']} route "
                    f"labels and was previously evaluated with {facts['train_rows']} "
                    f"training rows and {facts['test_rows']} held-out rows. Module 5 "
                    f"training produced an accuracy of {facts['accuracy']:.4f}, macro "
                    f"F1 score of {facts['macro_f1']:.4f}, and mean similarity of "
                    f"{facts['mean_similarity']:.4f}. The alpha interface then fits "
                    "the retriever across the full corpus and returns the matched route, "
                    "confidence score, and approved answer."
                ),
                (
                    f"The Module 6 demo includes {facts['demo_exchanges']} scripted "
                    f"customer messages. With a stricter alpha threshold of "
                    f"{facts['alpha_threshold']:.2f}, {facts['accepted']} messages were "
                    f"accepted as in-domain and {facts['fallback']} message was routed "
                    "to clarification or human escalation. That fallback behavior is "
                    "important because the Module 8 chatbot should avoid returning a "
                    "confident but irrelevant answer when the input is outside the "
                    "approved customer-service scope."
                ),
            ],
        ),
        (
            "Comparison With Final Goals",
            [
                (
                    "The alpha already does the core task that the final chatbot must "
                    "do: it responds to user text with route-specific support content "
                    "rather than nonsense. It is also closed-domain, which is the safer "
                    "choice for a business customer-service application because every "
                    "returned answer is tied to a controlled FAQ/SOP entry. This gives "
                    "the project a clearer audit trail than a broad open-domain chatbot."
                ),
                (
                    "The final version still needs several improvements. First, it "
                    "needs a more usable submission surface, such as a command-line "
                    "executable with clear instructions or a small local web chat. "
                    "Second, the response logic needs better confidence calibration so "
                    "the chatbot asks for clarification when the score is weak. Third, "
                    "the final report needs to document the tools, model type, open- or "
                    "closed-domain scope, and run instructions required by the Module 8 "
                    "portfolio prompt."
                ),
                (
                    "The alpha is therefore functional but not final. It proves the "
                    "retrieval design, but it does not yet prove that the chatbot can "
                    "handle messy real customer language, ambiguous requests, or "
                    "phrases that do not share obvious keywords with the FAQ/SOP "
                    "entries."
                ),
            ],
        ),
        (
            "Strengths and Gaps",
            [
                (
                    "The strongest part of the current project is control. The system "
                    "does not invent policies or attempt broad conversation. It maps "
                    "messages to approved answers, which is appropriate for a customer "
                    "service chatbot that could be adapted across companies. The design "
                    "is also reproducible because the corpus, thresholds, demo messages, "
                    "and output transcript are all included in the milestone package."
                ),
                (
                    "The main weakness is data coverage. The FAQ/SOP corpus is still "
                    "small and curated, so the high held-out score should be interpreted "
                    "as an alpha validation result rather than production evidence. The "
                    "model depends on shared vocabulary between the user message and "
                    "the approved entries. If customers use indirect language, emotional "
                    "framing, abbreviations, or unusual wording, the current unigram "
                    "retriever may rank the wrong route."
                ),
                (
                    "A second gap is threshold calibration. Module 5 used a draft "
                    f"similarity threshold of {facts['module5_threshold']:.2f}, but the "
                    "Module 6 alpha needed a stricter threshold to avoid over-answering "
                    "an out-of-domain request. That change is not a failure; it is the "
                    "kind of alpha evidence that should guide the final design."
                ),
            ],
        ),
        (
            "Improvement Strategy",
            [
                (
                    "The next step is retraining after the dataset is expanded. The "
                    "custom FAQ/SOP corpus should be enlarged with more paraphrases for "
                    "each route, more ambiguous messages, and explicit out-of-domain "
                    "examples. Retraining is necessary because the model cannot learn "
                    "customer phrasing that is not represented in the corpus."
                ),
                (
                    "The two public datasets selected in Module 4 should still be used "
                    "as supplemental validation assets. The RSiCS corpus can help test "
                    "whether the chatbot handles realistic customer-service phrasing "
                    "and relational language (Beaver et al., 2017). WikiQA can help "
                    "test answer ranking and no-answer behavior because it includes "
                    "question-answer relevance judgments (Yang et al., 2015)."
                ),
                (
                    "Technically, the final version can remain retrieval-based, but it "
                    "should evaluate additional feature settings such as bigrams, "
                    "character n-grams, and a calibrated fallback threshold. A later "
                    "implementation could also use a tested machine-learning library "
                    "for repeatable vectorization, model evaluation, and packaging "
                    "(Pedregosa et al., 2011). The key design principle should remain "
                    "the same: high-confidence matches return approved content, weak "
                    "matches ask for clarification, and unresolved cases escalate to a "
                    "human support queue."
                ),
            ],
        ),
        (
            "Conclusion",
            [
                (
                    "The Module 6 alpha is on track for the Module 8 portfolio project. "
                    "It demonstrates the core closed-domain chatbot workflow and "
                    "produces responses that are tied to customer-service input text. "
                    "The remaining work is practical rather than conceptual: expand and "
                    "stress-test the dataset, calibrate the fallback rule, improve the "
                    "runtime interface, and package the final chatbot with clear "
                    "instructions and source-supported explanation."
                ),
            ],
        ),
        (
            "Use of AI Assistance",
            [
                (
                    "AI assistance was used to help organize the alpha code, prepare "
                    "the milestone package, and draft this report. The chatbot script, "
                    "demo transcript, metrics, validation checks, DOCX, and PDF were "
                    "generated and reviewed locally in the CSC525 workspace before "
                    "submission."
                ),
            ],
        ),
    ]


def set_run_font(run, *, bold: bool = False, italic: bool = False) -> None:
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(12)
    run.bold = bold
    run.italic = italic


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()

    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char_separate = OxmlElement("w:fldChar")
    fld_char_separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")

    run._r.append(fld_char_begin)
    run._r.append(instr_text)
    run._r.append(fld_char_separate)
    run._r.append(text)
    run._r.append(fld_char_end)


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.5)
    section.footer_distance = Inches(0.5)

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 2.0
    normal.paragraph_format.space_after = Pt(0)

    for style_name in ("Heading 1", "Heading 2"):
        style = document.styles[style_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.size = Pt(12)
        style.font.bold = True
        style.paragraph_format.line_spacing = 2.0
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(0)

    add_page_number(section.header.paragraphs[0])


def add_centered_line(document: Document, text: str, *, bold: bool = False) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.line_spacing = 2.0
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    set_run_font(run, bold=bold)


def add_body_paragraph(document: Document, text: str, *, indent: bool = True) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.line_spacing = 2.0
    paragraph.paragraph_format.space_after = Pt(0)
    if indent:
        paragraph.paragraph_format.first_line_indent = Inches(0.5)
    run = paragraph.add_run(text)
    set_run_font(run)


def add_section_heading(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.line_spacing = 2.0
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(text)
    set_run_font(run, bold=True)


def add_reference(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.line_spacing = 2.0
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.left_indent = Inches(0.5)
    paragraph.paragraph_format.first_line_indent = Inches(-0.5)
    run = paragraph.add_run(text)
    set_run_font(run)


def build_docx(facts: dict[str, Any]) -> None:
    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    document = Document()
    configure_document(document)

    for _ in range(4):
        document.add_paragraph()
    add_centered_line(document, TITLE, bold=True)
    document.add_paragraph()
    for line in (STUDENT_NAME, UNIVERSITY, COURSE, ASSIGNMENT, INSTRUCTOR, DATE):
        add_centered_line(document, line)

    document.add_page_break()
    add_centered_line(document, TITLE, bold=True)
    for heading, paragraphs in body_sections(facts):
        if heading is not None:
            if heading in {"Strengths and Gaps"}:
                document.add_page_break()
            add_section_heading(document, heading)
        for paragraph in paragraphs:
            add_body_paragraph(document, paragraph)

    document.add_page_break()
    add_centered_line(document, "References", bold=True)
    for reference in REFERENCES:
        add_reference(document, reference)

    document.save(DOCX_PATH)
    print(DOCX_PATH)


def pdf_paragraph(text: str, style: ParagraphStyle) -> Paragraph:
    return Paragraph(escape(text), style)


def on_page(canvas, _doc) -> None:
    canvas.saveState()
    canvas.setFont("Times-Roman", 12)
    canvas.drawRightString(7.5 * inch, 10.5 * inch, str(canvas.getPageNumber()))
    canvas.restoreState()


def build_pdf(facts: dict[str, Any]) -> None:
    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    document = SimpleDocTemplate(
        str(PDF_PATH),
        pagesize=letter,
        rightMargin=inch,
        leftMargin=inch,
        topMargin=inch,
        bottomMargin=inch,
    )

    title_style = ParagraphStyle(
        "TitleStyle",
        fontName="Times-Bold",
        fontSize=12,
        leading=24,
        alignment=TA_CENTER,
        spaceAfter=0,
    )
    center_style = ParagraphStyle(
        "CenterStyle",
        fontName="Times-Roman",
        fontSize=12,
        leading=24,
        alignment=TA_CENTER,
        spaceAfter=0,
    )
    body_style = ParagraphStyle(
        "BodyStyle",
        fontName="Times-Roman",
        fontSize=12,
        leading=24,
        firstLineIndent=0.5 * inch,
        alignment=TA_LEFT,
        spaceAfter=0,
    )
    heading_style = ParagraphStyle(
        "HeadingStyle",
        fontName="Times-Bold",
        fontSize=12,
        leading=24,
        alignment=TA_LEFT,
        spaceAfter=0,
    )
    refs_style = ParagraphStyle(
        "ReferencesStyle",
        fontName="Times-Roman",
        fontSize=12,
        leading=24,
        leftIndent=0.5 * inch,
        firstLineIndent=-0.5 * inch,
        alignment=TA_LEFT,
        spaceAfter=0,
    )

    story = [
        Spacer(1, 1.35 * inch),
        pdf_paragraph(TITLE, title_style),
        Spacer(1, 0.24 * inch),
    ]
    for line in (STUDENT_NAME, UNIVERSITY, COURSE, ASSIGNMENT, INSTRUCTOR, DATE):
        story.append(pdf_paragraph(line, center_style))

    story.append(PageBreak())
    story.append(pdf_paragraph(TITLE, title_style))
    for heading, paragraphs in body_sections(facts):
        if heading is not None:
            if heading in {"Strengths and Gaps"}:
                story.append(PageBreak())
            story.append(pdf_paragraph(heading, heading_style))
        for paragraph in paragraphs:
            story.append(pdf_paragraph(paragraph, body_style))

    story.append(PageBreak())
    story.append(pdf_paragraph("References", title_style))
    for reference in REFERENCES:
        story.append(pdf_paragraph(reference, refs_style))

    document.build(story, onFirstPage=on_page, onLaterPages=on_page)
    print(PDF_PATH)


def main() -> None:
    metrics = load_metrics()
    predictions = load_demo_predictions()
    facts = project_facts(metrics, predictions)
    build_docx(facts)
    build_pdf(facts)


if __name__ == "__main__":
    main()
