"""Build the APA-style CSC580 Module 3 Option 2 Word analysis."""

from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "outputs"
DOCX_PATH = ROOT / "CSC580_CTA_3_2_Dunn_Justin.docx"


def set_cell_margins(cell, top=100, start=100, bottom=100, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr, fld_char2])


def add_caption(doc: Document, number: int, title: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.keep_with_next = True
    r1 = p.add_run(f"Figure {number}\n")
    r1.bold = True
    r2 = p.add_run(title)
    r2.italic = True


def add_figure(doc: Document, filename: str, number: int, title: str, width: float = 6.35) -> None:
    add_caption(doc, number, title)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(3)
    p.add_run().add_picture(str(OUT / filename), width=Inches(width))


def add_body_paragraph(doc: Document, text: str) -> None:
    p = doc.add_paragraph(text)
    p.paragraph_format.first_line_indent = Inches(0.5)


def main() -> None:
    results = json.loads((OUT / "results.json").read_text(encoding="utf-8"))
    mse = results["mse_loss_model"]
    mae = results["mae_loss_model"]

    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(1)
    sec.bottom_margin = Inches(1)
    sec.left_margin = Inches(1)
    sec.right_margin = Inches(1)
    add_page_number(sec.header.paragraphs[0])

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.paragraph_format.line_spacing = 2.0
    normal.paragraph_format.space_after = Pt(0)
    for style_name in ["Heading 1", "Heading 2"]:
        style = styles[style_name]
        style.font.name = "Calibri"
        style.font.color.rgb = None
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    styles["Heading 1"].font.size = Pt(11)
    styles["Heading 1"].font.bold = True
    styles["Heading 1"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    styles["Heading 2"].font.size = Pt(11)
    styles["Heading 2"].font.bold = True

    # APA student title page.
    for _ in range(4):
        doc.add_paragraph()
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Option 2: Predicting Fuel Efficiency Using TensorFlow").bold = True
    for line in [
        "Justin Dunn",
        "Colorado State University Global",
        "CSC580-1: Applying Machine Learning and Neural Networks",
        "Dong Nguyen",
        "August 9, 2026",
    ]:
        p = doc.add_paragraph(line)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_page_break()

    heading = doc.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading.add_run("Option 2: Predicting Fuel Efficiency Using TensorFlow").bold = True

    doc.add_heading("Introduction", level=1)
    add_body_paragraph(
        doc,
        "This analysis developed and evaluated fully connected TensorFlow regression models that predict "
        "automobile fuel efficiency in miles per gallon (MPG). The work used the Auto MPG dataset, which "
        "contains vehicle characteristics from the late 1970s and early 1980s (UCI Machine Learning "
        "Repository, 1993). The assignment objective was to complete the full regression workflow: inspect "
        "and prepare the data, normalize differently scaled features, build a neural network, train it for "
        "1,000 epochs, visualize learning behavior, and compare mean absolute error (MAE) with mean squared "
        "error (MSE). The central conclusion is that both models were useful, but the MAE-loss model "
        "generalized better on the held-out test set in this reproducible run."
    )

    doc.add_heading("Data Preparation and Exploration", level=1)
    add_body_paragraph(
        doc,
        f"The original data were downloaded programmatically and parsed with the eight assignment variables. "
        f"Rows with missing horsepower were removed, leaving {results['rows_after_cleaning']} complete records. "
        "Origin was converted into separate USA, Europe, and Japan indicator variables so that the network "
        "would not treat the geographic codes as a continuous numerical scale. The cleaned data were split "
        f"into {results['training_rows']} training records and {results['test_rows']} test records. Means and "
        "standard deviations were calculated from the training features only, and both partitions were "
        "normalized with those training statistics. This ordering prevents information from the test set from "
        "leaking into model development."
    )
    add_figure(doc, "01_dataset_tail.png", 1, "Tail of the Cleaned and Encoded Auto MPG Dataset", 6.25)
    add_body_paragraph(
        doc,
        "Figure 2 shows that MPG has strong inverse relationships with displacement and weight. Cylinders "
        "also form distinct bands, reflecting its discrete nature. These relationships support a regression "
        "model because vehicle efficiency varies systematically with engine size and mass, while the curved "
        "and clustered patterns justify a nonlinear network rather than assuming one straight-line effect for "
        "every predictor."
    )
    add_figure(doc, "02_pairplot.png", 2, "Pairwise Relationships Among MPG, Cylinders, Displacement, and Weight", 6.0)
    add_figure(doc, "03_training_statistics.png", 3, "Descriptive Statistics Used for Feature Normalization", 6.25)

    doc.add_heading("Model Architecture and Training", level=1)
    add_body_paragraph(
        doc,
        "The primary model used the TensorFlow Keras Sequential API (TensorFlow, 2026). It contained nine "
        "normalized inputs, two densely connected hidden layers with 64 rectified linear unit neurons each, "
        "and one linear output neuron for the continuous MPG prediction. This architecture contained 4,865 "
        "trainable parameters. RMSprop optimization used a learning rate of 0.001. The primary model minimized "
        "MSE while recording both MAE and MSE. A second model used the same architecture, initialization seed, "
        "optimizer, split, and 1,000-epoch budget but minimized MAE, allowing the loss functions to be compared "
        "without changing the broader design."
    )
    add_figure(doc, "04_model_summary.png", 4, "TensorFlow Keras Model Architecture and Parameter Counts", 6.0)
    add_body_paragraph(
        doc,
        "Before training, predictions were near zero and therefore far from the observed MPG values. This "
        "baseline confirms that the final estimates resulted from learning rather than an untrained network "
        "accidentally producing plausible outputs."
    )
    add_figure(doc, "05_untrained_predictions.png", 5, "Ten Predictions Produced Before Model Training", 5.8)

    doc.add_heading("Training Progress", level=1)
    add_body_paragraph(
        doc,
        f"The MSE-loss model completed all 1,000 epochs. Its lowest validation MAE was "
        f"{results['mse_model_best_validation_mae']:.3f} MPG, and its lowest validation MSE was "
        f"{results['mse_model_best_validation_mse']:.3f} MPG squared. The steep initial improvement in Figures "
        "7 and 8 indicates that the network quickly learned the dominant relationships. Training error "
        "continued to fall after validation error flattened, creating a widening train-validation gap. This "
        "is evidence of overfitting during later epochs and suggests that early stopping would be an appropriate "
        "production improvement, although all 1,000 epochs were retained to satisfy the assignment."
    )
    add_figure(doc, "06_history_tail.png", 6, "Final Five Rows of the 1,000-Epoch Training History", 6.2)
    add_figure(doc, "07_history_mae.png", 7, "Training and Validation Mean Absolute Error", 6.2)
    add_figure(doc, "08_history_mse.png", 8, "Training and Validation Mean Squared Error", 6.2)

    doc.add_heading("Model Evaluation and Comparison", level=1)
    add_body_paragraph(
        doc,
        f"On the untouched test set, the MSE-loss model obtained an MAE of {mse['mae']:.3f} MPG, an MSE of "
        f"{mse['mse']:.3f} MPG squared, and an RMSE of {results['mse_loss_model_rmse']:.3f} MPG. The MAE-loss "
        f"model improved all three evaluation measures: MAE was {mae['mae']:.3f} MPG, MSE was "
        f"{mae['mse']:.3f} MPG squared, and RMSE was {results['mae_loss_model_rmse']:.3f} MPG. Therefore, the "
        "MAE-loss model was the better fitted model for this split. MAE assigns a linear penalty to each error "
        "and is less dominated by a few unusually large residuals, whereas MSE squares errors and emphasizes "
        "large misses. MSE remains useful when large fuel-efficiency errors carry disproportionate costs, but "
        "the lower held-out MAE and RMSE make the MAE-loss model the stronger general-purpose choice here."
    )
    add_figure(doc, "09_model_comparison.png", 9, "Held-Out Test Performance for MAE-Loss and MSE-Loss Models", 6.0)
    add_body_paragraph(
        doc,
        f"Figure 10 shows that the primary model tracked the perfect-prediction line across most of the MPG "
        f"range, although dispersion increased among higher-efficiency vehicles. The mean residual was "
        f"{results['mse_model_mean_residual']:.3f} MPG, which is close to zero and indicates little aggregate "
        "directional bias. The residual standard deviation was "
        f"{results['mse_model_residual_std']:.3f} MPG. Figure 11 nevertheless shows several errors beyond plus "
        "or minus 5 MPG; these observations explain why RMSE exceeded MAE and should be investigated before "
        "using the model for high-stakes decisions."
    )
    add_figure(doc, "10_predicted_vs_actual.png", 10, "Predicted MPG Compared With True MPG for the MSE-Loss Model", 5.7)
    add_figure(doc, "11_residual_distribution.png", 11, "Distribution of Test-Set Prediction Errors", 6.1)

    doc.add_heading("Conclusion", level=1)
    add_body_paragraph(
        doc,
        "The completed workflow demonstrated that a compact fully connected neural network can make useful "
        "fuel-efficiency predictions from historical vehicle specifications. Careful cleaning, one-hot "
        "encoding, train-only normalization, and a held-out test set made the comparison defensible. Both loss "
        "functions produced practical models, but the MAE-loss network achieved the lower test MAE, MSE, and "
        "RMSE and was therefore selected as the better fitted model. The learning curves also showed that "
        "training for the entire 1,000-epoch requirement exceeded the point of best validation performance. A "
        "next iteration should restore the best validation weights through early stopping, repeat evaluation "
        "across multiple random splits, and analyze the largest residuals by model year and origin."
    )

    doc.add_page_break()
    doc.add_heading("References", level=1)
    refs = [
        "TensorFlow. (2026). *tf.keras API*. https://www.tensorflow.org/api_docs/python/tf/keras",
        "UCI Machine Learning Repository. (1993). *Auto MPG* [Data set]. https://doi.org/10.24432/C5859H",
    ]
    for ref in refs:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.5)
        p.paragraph_format.first_line_indent = Inches(-0.5)
        # Minimal markdown-style italics are converted into italic runs.
        parts = ref.split("*")
        for idx, part in enumerate(parts):
            run = p.add_run(part)
            run.italic = idx % 2 == 1

    doc.core_properties.title = "Option 2: Predicting Fuel Efficiency Using TensorFlow"
    doc.core_properties.author = "Justin Dunn"
    doc.core_properties.subject = "CSC580 Module 3 Critical Thinking Assignment"
    doc.save(DOCX_PATH)
    print(DOCX_PATH)


if __name__ == "__main__":
    main()
